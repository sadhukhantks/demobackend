"""
Sales Report endpoints.

Endpoints for Tally-style downloads and other sales reports. Mirrors the
validation / response shape of src/juteProcurement/reports.py.
"""

import logging
import re
from collections import defaultdict
from datetime import datetime
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy.orm import Session

from src.authorization.utils import get_current_user_with_refresh
from src.config.db import get_tenant_db
from src.sales.constants import INVOICE_TYPE_IDS, SALES_STATUS_IDS
from src.sales.reportQueries import (
    get_invoice_list_report_query,
    get_sales_jute_summary_query,
    get_sales_jute_tally_header_query,
    get_sales_jute_tally_line_query,
    get_sales_order_list_report_query,
)

logger = logging.getLogger(__name__)

router = APIRouter()

# Hard cap on rows returned by the tally download endpoint. Matches the
# 10,000-row limit used elsewhere in the reports module.
_MAX_ROWS = 10000

# Tally "Accounting Voucher" import columns, keyed by docs/Sales.xlsx column
# letter. Header strings are copied VERBATIM from that sheet — lowercase AA and
# the typos at AG / AJ are what the Tally import template matches on. Do not
# "fix" them. Exception: AH is "Voucher Narration" (the sheet said "Ledger
# Narration"; user wants the voucher-level narration). Dropped as duplicates:
# O "Change Mode " (never filled; AA carries it), U "New Ref No." (voucher no
# + lot again) and Y "Delivery Note Date/..." (same challan text as W) — so
# from P on the letter is a dict key only and the sheet column is the dict
# position.
_TALLY_COLUMNS = {
    "A": "Voucher Date",
    "B": "Voucher Type Name",
    "C": "Voucher Number",
    "D": "Buyer/Supplier - Address",
    "E": "Buyer/Supplier - Pincode",
    "F": "Ledger Name",
    "G": "Ledger Amount",
    "H": "Ledger Amount Dr/Cr",
    "I": "Item Name",
    "J": "item description",
    "K": "Billed Quantity",
    "L": "Item Rate",
    "M": "Item Rate per",
    "N": "Item Amount",
    "P": "Buyer/Supplier - Bill to/from",
    "Q": "Dispatched through",
    "R": "Destination",
    "S": "Reference No.",
    "T": "batch lot no",
    "V": "Godown",
    "W": "Despatch Doc No",
    "X": "Consignee (ship to)",
    "Z": "Description of Ledger",
    "AA": "change mode",
    "AB": "order no",
    "AC": "order dt",
    "AD": "terms of payment",
    "AE": "terms of delivery",
    "AF": "state",
    "AG": "ship to adrress",
    "AH": "Voucher Narration",
    "AI": "Place of supply",
    "AJ": "state consginee",
    "AK": "consignee zip",
}

_TALLY_NUMBER_FORMATS = {
    "A": "dd-mm-yyyy",
    "G": "0.00",
    "K": "#,##0",
    "L": "0.00",
    "N": "0.00",
}


def _parse_iso_date(value: str, field_name: str) -> str:
    """Validate a YYYY-MM-DD date string. Raises HTTPException(400) if invalid."""
    try:
        datetime.strptime(value, "%Y-%m-%d")
    except (TypeError, ValueError):
        raise HTTPException(
            status_code=400,
            detail=f"Invalid {field_name} format, expected YYYY-MM-DD",
        )
    return value


def _dmy(d) -> str:
    return d.strftime("%d.%m.%Y") if d else ""


def _same_place(a, b) -> bool:
    """Dispatch-from and ship-to are the same address, ignoring case,
    spacing and punctuation. Blank on either side is never "same"."""
    norm = lambda s: re.sub(r"[^a-z0-9]", "", (s or "").lower())
    return bool(norm(a)) and norm(a) == norm(b)


def _kg_text(kg: float) -> str:
    """24073.0 -> '24073', 51.5 -> '51.5'."""
    return f"{kg:.3f}".rstrip("0").rstrip(".")


def _line_kg(line: dict):
    """(qty_kg, rate_per_kg) for one sales_invoice_dtl row."""
    qty = float(line["quantity"] or 0)
    rate = float(line["rate"] or 0)
    # ponytail: legacy (pre-2026) raw-jute lines are quintals - uom QNT, or a
    # NULL uom at quintal-scale rates (~5000-15000/qtl) on dev3 + sls. Ceiling:
    # a future kg line saved with a NULL uom would be read as quintals; move to
    # an explicit uom -> factor map if that ever shows up.
    if line["uom_id"] is None or line["uom_name"] == "QNT":
        return round(qty * 100, 3), round(rate / 100, 4)
    return qty, rate


def build_tally_voucher_rows(headers, lines) -> list:
    """Build Tally "Accounting Voucher" rows (dicts keyed by column letter).

    headers: one dict per invoice (get_sales_jute_tally_header_query).
    lines:   one dict per invoice line (get_sales_jute_tally_line_query).

    Per invoice with lines: party row, sales ledger, one item row per line with
    qty > 0, claim (claim > 0), round off (non-zero) -- ledger names are the
    company's (header *_ledger columns, tally_link_mst). All
    money comes from the lines (sales_invoice.invoice_amount is unreliable) and
    is rounded to 2dp first; the party amount is derived from the rounded
    parts, so every voucher's Ledger Amounts sum to exactly 0.
    """
    lines_by_invoice = defaultdict(list)
    for line in lines:
        lines_by_invoice[line["invoice_id"]].append(line)

    rows = []
    for h in headers:
        inv_lines = lines_by_invoice.get(h["invoice_id"])
        if not inv_lines:
            continue  # no lines = nothing to bill (bad data); never a 0-value voucher
        vno = h["voucher_no"]
        lot = h["lot_no"] or ""
        party = h["party_name"]
        base = {"A": h["invoice_date"], "B": "Sales", "C": vno, "AA": "Item Invoice"}

        items = round(sum(round(float(ln["amount_without_tax"] or 0), 2) for ln in inv_lines), 2)
        claim = round(sum(float(ln["claim_amount"] or 0) for ln in inv_lines), 2)
        ro = round(float(h["round_off"] or 0), 2)
        party_amt = round(items - claim + ro, 2)

        is_bale = (h["unit_conversion"] or "").upper() == "BALE"
        item_rows, total_kg, claim_kg_by_rate = [], 0.0, {}
        for ln in inv_lines:
            kg, rate = _line_kg(ln)
            if kg <= 0:
                continue
            total_kg += kg
            line_claim = float(ln["claim_amount"] or 0)
            if line_claim > 0:
                # claim_rate units vary by era; derive the per-kg rate instead.
                claim_rate = round(line_claim / kg, 2)
                claim_kg_by_rate[claim_rate] = claim_kg_by_rate.get(claim_rate, 0) + kg
            item_rows.append({
                **base,
                "I": ln["item_name"],
                "J": f"Raw Jute: {ln['bales']} Bales" if is_bale and ln["bales"] else "Raw Jute: Loose",
                "K": kg,
                "L": rate,
                "M": "kg",
                "N": round(float(ln["amount_without_tax"] or 0), 2),
                "S": f"REF NO: {lot}" if lot else "",
                "T": lot,
                "V": h["godown_name"],
            })

        challan = (h["challan_no"] or "").strip()
        challan_text = ""
        if challan:
            challan_text = f"Ch.No: {challan}"
            if h["challan_date"]:
                challan_text += f" , Dt: {_dmy(h['challan_date'])}"
        ship_state = (h["ship_state"] or "").title()
        ref_text = f", REF NO: {lot}" if lot else ""
        amount_text = f"{abs(party_amt):.2f}".removesuffix(".00")
        rows.append({
            **base,
            "D": h["bill_address"] or "",
            "E": h["bill_zip"],
            "F": party,
            "G": -party_amt,
            # No truck when the goods never moved (dispatch-from == ship-to).
            "Q": (f"Truck No: {h['vehicle_no']}"
                  if h["vehicle_no"] and not _same_place(h["dispatch_address"], h["ship_address"])
                  else ""),
            "R": h["destination"] or "",   # fixed per receiving mill / buyer, see header query
            "W": challan_text,
            "X": h["consignee_name"],
            "AB": h["buyer_order_no"] or "",
            "AC": _dmy(h["buyer_order_date"]),
            "AD": f"{h['payment_terms']} days" if h["payment_terms"] else "",
            "AF": (h["bill_state"] or "").title(),
            "AG": h["ship_address"] or "",
            "AH": (
                f"Being {_kg_text(total_kg)} Kg. Raw Jute Sold To '{party}' "
                f"against Inv No: {vno}, Dt: {_dmy(h['invoice_date'])}{ref_text}, "
                f"Amount Rs. {amount_text}/-"
            ),
            "AI": ship_state,
            "AJ": ship_state,
            "AK": h["ship_zip"],
        })
        rows.append({**base, "F": h["sales_ledger"], "G": items})
        rows.extend(item_rows)
        if claim > 0:
            rows.append({
                **base,
                "F": h["claim_ledger"],
                "G": -claim,
                "Z": " + ".join(
                    f"(Q): {_kg_text(kg)} Kg. X Rs.{rate:.2f}/-"
                    for rate, kg in claim_kg_by_rate.items()
                ),
            })
        if ro:
            rows.append({**base, "F": h["round_off_ledger"], "G": ro})
    return rows


def _build_tally_xlsx(rows) -> BytesIO:
    """Single-sheet 'Accounting Voucher' workbook for the Tally import."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Accounting Voucher"
    ws.append(list(_TALLY_COLUMNS.values()))
    for row in rows:
        ws.append([None if row.get(c) == "" else row.get(c) for c in _TALLY_COLUMNS])
        for col, fmt in _TALLY_NUMBER_FORMATS.items():
            if row.get(col) not in (None, ""):
                ws[f"{col}{ws.max_row}"].number_format = fmt
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


@router.get("/jute-tally-download")
def get_sales_jute_tally_download(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """Sales Raw Jute Tally xlsx download.

    Returns a single-sheet 'Accounting Voucher' workbook in the Tally import
    layout of docs/Sales.xlsx for approved Raw Jute invoices between two
    dates — one balanced voucher per invoice (see build_tally_voucher_rows).

    Query params:
    - co_id:     Company ID (required, int)
    - branch_id: Branch ID (optional, int)
    - date_from: Start date YYYY-MM-DD (required)
    - date_to:   End date YYYY-MM-DD (required)

    Response: `application/vnd.openxmlformats-officedocument.spreadsheetml.sheet`
    with Content-Disposition attachment.
    """
    try:
        q_co_id = request.query_params.get("co_id")
        q_branch_id = request.query_params.get("branch_id")
        q_date_from = request.query_params.get("date_from")
        q_date_to = request.query_params.get("date_to")

        if not q_co_id:
            raise HTTPException(status_code=400, detail="co_id is required")
        if not q_date_from:
            raise HTTPException(status_code=400, detail="date_from is required")
        if not q_date_to:
            raise HTTPException(status_code=400, detail="date_to is required")

        try:
            co_id = int(q_co_id)
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Invalid co_id")

        branch_id = None
        if q_branch_id not in (None, ""):
            try:
                branch_id = int(q_branch_id)
            except (TypeError, ValueError):
                raise HTTPException(status_code=400, detail="Invalid branch_id")

        date_from = _parse_iso_date(q_date_from, "date_from")
        date_to = _parse_iso_date(q_date_to, "date_to")

        params = {
            "co_id": co_id,
            "branch_id": branch_id,
            "date_from": date_from,
            "date_to": date_to,
            "invoice_type_id": INVOICE_TYPE_IDS["RAW_JUTE"],
            "approved_status_id": SALES_STATUS_IDS["APPROVED"],
        }

        headers = [
            dict(r._mapping)
            for r in db.execute(get_sales_jute_tally_header_query(), params).fetchall()
        ]
        lines = [
            dict(r._mapping)
            for r in db.execute(get_sales_jute_tally_line_query(), params).fetchall()
        ]
        rows = build_tally_voucher_rows(headers, lines)

        if len(rows) > _MAX_ROWS:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Result set too large ({len(rows)} rows, "
                    f"max {_MAX_ROWS}). Narrow your date range."
                ),
            )

        bio = _build_tally_xlsx(rows)

        filename = f"sales_raw_jute_tally_{date_from}_{date_to}.xlsx"
        return StreamingResponse(
            bio,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
            },
        )

    except HTTPException:
        raise
    except Exception as e:
        logger.error(
            f"Error generating sales jute tally download: {e}", exc_info=True
        )
        raise HTTPException(
            status_code=500,
            detail=f"Error generating sales jute tally download: {str(e)}",
        )


@router.get("/jute-mr-summary")
def get_sales_jute_mr_summary(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """Per-sales-invoice summary of MRs that match the Tally download scope.

    Same filter scope as `/jute-tally-download`. One row per sales invoice
    with linked MR number + amounts so the user can preview before downloading.

    Query params:
    - co_id:     Company ID (required, int)
    - branch_id: Branch ID (optional, int)
    - date_from: Start date YYYY-MM-DD (required)
    - date_to:   End date YYYY-MM-DD (required)

    Returns: {"data": [row, ...]}
    """
    try:
        q_co_id = request.query_params.get("co_id")
        q_branch_id = request.query_params.get("branch_id")
        q_date_from = request.query_params.get("date_from")
        q_date_to = request.query_params.get("date_to")

        if not q_co_id:
            raise HTTPException(status_code=400, detail="co_id is required")
        if not q_date_from:
            raise HTTPException(status_code=400, detail="date_from is required")
        if not q_date_to:
            raise HTTPException(status_code=400, detail="date_to is required")

        try:
            co_id = int(q_co_id)
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Invalid co_id")

        branch_id = None
        if q_branch_id not in (None, ""):
            try:
                branch_id = int(q_branch_id)
            except (TypeError, ValueError):
                raise HTTPException(status_code=400, detail="Invalid branch_id")

        date_from = _parse_iso_date(q_date_from, "date_from")
        date_to = _parse_iso_date(q_date_to, "date_to")

        params = {
            "co_id": co_id,
            "branch_id": branch_id,
            "date_from": date_from,
            "date_to": date_to,
            "invoice_type_id": INVOICE_TYPE_IDS["RAW_JUTE"],
            "approved_status_id": SALES_STATUS_IDS["APPROVED"],
        }

        rows = db.execute(get_sales_jute_summary_query(), params).fetchall()

        data = []
        for row in rows:
            mapped = dict(row._mapping)
            value = mapped.get("invoice_date")
            if value is not None and hasattr(value, "isoformat"):
                mapped["invoice_date"] = value.isoformat()
            for num_field in ("total_qty_kg", "invoice_amount", "claim_amount"):
                v = mapped.get(num_field)
                if v is not None:
                    mapped[num_field] = float(v)
            data.append(mapped)

        return {"data": data}

    except HTTPException:
        raise
    except Exception as e:
        logger.error(
            f"Error fetching sales jute MR summary: {e}", exc_info=True
        )
        raise HTTPException(
            status_code=500,
            detail=f"Error fetching sales jute MR summary: {str(e)}",
        )


# =============================================================================
# Sales Order list / Invoice list reports
#
# Two simple, date-range scoped reports surfaced under the Sales Reports page.
# All statuses included (incl. Cancelled) per user requirement — Status column
# is exposed so users can scan visually. JSON variant powers the on-screen
# DataGrid; -download variant streams an xlsx via openpyxl.
# =============================================================================

_SO_LIST_DATE_COLS = ("sales_order_date",)
_SO_LIST_NUM_COLS = ("gross_amount", "net_amount")

_INV_LIST_DATE_COLS = ("invoice_date", "challan_date", "due_date")
_INV_LIST_NUM_COLS = (
    "invoice_amount",
    "tax_amount",
    "tax_payable",
    "round_off",
)


def _serialize_report_row(mapped: dict, date_cols: tuple, num_cols: tuple) -> dict:
    """Convert SQL date/numeric columns to JSON-friendly types."""
    for col in date_cols:
        v = mapped.get(col)
        if v is not None and hasattr(v, "isoformat"):
            mapped[col] = v.isoformat()
    for col in num_cols:
        v = mapped.get(col)
        if v is not None:
            try:
                mapped[col] = float(v)
            except (TypeError, ValueError):
                pass
    return mapped


def _read_list_query_params(request: Request) -> dict:
    """Validate the shared (co_id, branch_id, date_from, date_to) tuple."""
    q_co_id = request.query_params.get("co_id")
    q_branch_id = request.query_params.get("branch_id")
    q_date_from = request.query_params.get("date_from")
    q_date_to = request.query_params.get("date_to")

    if not q_co_id:
        raise HTTPException(status_code=400, detail="co_id is required")
    if not q_date_from:
        raise HTTPException(status_code=400, detail="date_from is required")
    if not q_date_to:
        raise HTTPException(status_code=400, detail="date_to is required")

    try:
        co_id = int(q_co_id)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="Invalid co_id")

    branch_id = None
    if q_branch_id not in (None, ""):
        try:
            branch_id = int(q_branch_id)
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Invalid branch_id")

    return {
        "co_id": co_id,
        "branch_id": branch_id,
        "date_from": _parse_iso_date(q_date_from, "date_from"),
        "date_to": _parse_iso_date(q_date_to, "date_to"),
    }


_SO_XLSX_COLUMNS = [
    ("sales_no", "Sales No"),
    ("sales_no_formatted", "SO Number"),
    ("sales_order_date", "Date"),
    ("branch_name", "Branch"),
    ("party_name", "Party"),
    ("quotation_no", "Quotation No"),
    ("invoice_type", "Invoice Type"),
    ("status_name", "Status"),
    ("approval_level", "Approval Level"),
    ("gross_amount", "Gross Amount"),
    ("net_amount", "Net Amount"),
]

_INV_XLSX_COLUMNS = [
    ("invoice_no", "Invoice No"),
    ("invoice_no_formatted", "Invoice Number"),
    ("invoice_date", "Date"),
    ("branch_name", "Branch"),
    ("party_name", "Party"),
    ("invoice_type", "Invoice Type"),
    ("status_name", "Status"),
    ("approval_level", "Approval Level"),
    ("challan_no", "Challan No"),
    ("challan_date", "Challan Date"),
    ("due_date", "Due Date"),
    ("invoice_amount", "Invoice Amount"),
    ("tax_amount", "Tax Amount"),
    ("tax_payable", "Tax Payable"),
    ("round_off", "Round Off"),
]


def _build_list_xlsx(rows, columns, sheet_title: str) -> BytesIO:
    """Build a single-sheet xlsx with the given columns from row mappings."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    ws.append([label for _, label in columns])
    for r in rows:
        mapped = dict(r._mapping)
        out_row = []
        for key, _ in columns:
            v = mapped.get(key)
            if v is not None and hasattr(v, "isoformat"):
                v = v.isoformat()
            out_row.append(v if v is not None else "")
        ws.append(out_row)

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


@router.get("/sales-order-list")
def get_sales_order_list(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """List of Sales Orders within a date range, JSON for on-screen render."""
    try:
        params = _read_list_query_params(request)
        rows = db.execute(get_sales_order_list_report_query(), params).fetchall()

        data = [
            _serialize_report_row(
                dict(r._mapping), _SO_LIST_DATE_COLS, _SO_LIST_NUM_COLS
            )
            for r in rows
        ]
        if len(data) > _MAX_ROWS:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Result set too large ({len(data)} rows, "
                    f"max {_MAX_ROWS}). Narrow your date range."
                ),
            )
        return {"data": data}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching sales order list report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail=f"Error fetching sales order list report: {str(e)}",
        )


@router.get("/sales-order-list-download")
def get_sales_order_list_download(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """Excel download of the Sales Orders list report."""
    try:
        params = _read_list_query_params(request)
        rows = db.execute(get_sales_order_list_report_query(), params).fetchall()
        if len(rows) > _MAX_ROWS:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Result set too large ({len(rows)} rows, "
                    f"max {_MAX_ROWS}). Narrow your date range."
                ),
            )
        bio = _build_list_xlsx(rows, _SO_XLSX_COLUMNS, "SalesOrders")
        filename = f"sales_orders_{params['date_from']}_{params['date_to']}.xlsx"
        return StreamingResponse(
            bio,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    except HTTPException:
        raise
    except Exception as e:
        logger.error(
            f"Error generating sales order list xlsx: {e}", exc_info=True
        )
        raise HTTPException(
            status_code=500,
            detail=f"Error generating sales order list xlsx: {str(e)}",
        )


@router.get("/invoice-list")
def get_invoice_list(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """List of Sales Invoices within a date range, JSON for on-screen render."""
    try:
        params = _read_list_query_params(request)
        rows = db.execute(get_invoice_list_report_query(), params).fetchall()

        data = [
            _serialize_report_row(
                dict(r._mapping), _INV_LIST_DATE_COLS, _INV_LIST_NUM_COLS
            )
            for r in rows
        ]
        if len(data) > _MAX_ROWS:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Result set too large ({len(data)} rows, "
                    f"max {_MAX_ROWS}). Narrow your date range."
                ),
            )
        return {"data": data}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching invoice list report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail=f"Error fetching invoice list report: {str(e)}",
        )


@router.get("/invoice-list-download")
def get_invoice_list_download(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """Excel download of the Sales Invoices list report."""
    try:
        params = _read_list_query_params(request)
        rows = db.execute(get_invoice_list_report_query(), params).fetchall()
        if len(rows) > _MAX_ROWS:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Result set too large ({len(rows)} rows, "
                    f"max {_MAX_ROWS}). Narrow your date range."
                ),
            )
        bio = _build_list_xlsx(rows, _INV_XLSX_COLUMNS, "Invoices")
        filename = f"sales_invoices_{params['date_from']}_{params['date_to']}.xlsx"
        return StreamingResponse(
            bio,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    except HTTPException:
        raise
    except Exception as e:
        logger.error(
            f"Error generating invoice list xlsx: {e}", exc_info=True
        )
        raise HTTPException(
            status_code=500,
            detail=f"Error generating invoice list xlsx: {str(e)}",
        )
