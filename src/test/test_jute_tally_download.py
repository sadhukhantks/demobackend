"""
Tests for the Jute Purchase Tally download (GET /api/juteReports/tally-download).

The export fills docs/PurchaseFormat.xlsx (Tally Prime "Accounting Voucher"):
one balanced voucher per approved MR, built by build_purchase_tally_rows.
"""

from datetime import date
from io import BytesIO
from unittest.mock import MagicMock

import openpyxl
import pytest
from fastapi.testclient import TestClient

from src.authorization.utils import get_current_user_with_refresh
from src.config.db import get_tenant_db
from src.juteProcurement.reports import (
    _PURCHASE_TALLY_TEMPLATE,
    build_purchase_tally_rows,
)
from src.main import app

client = TestClient(app)


def _header(**overrides):
    """One MR, shaped like get_jute_purchase_tally_header_query."""
    h = {
        "jute_mr_id": 1,
        "mr_no": "ACPL/JMR/26-27/00024",
        "jute_mr_date": date(2026, 6, 16),
        "branch_mr_no": 24,
        "party_id": "8668",
        "invoice_no": "SCP/MAY/021/26-27",
        "invoice_date": date(2026, 5, 23),
        "challan_no": "MAY/021",
        "challan_date": date(2026, 5, 23),
        "unit_conversion": "BALE",
        "tds_amount": 694.0,
        "party_name": "Snehjyoti Commerce Pvt Ltd (J-P)",
        "party_address": "7ram Kumar Rakhit Lane, Ground Floor, Kolkata",
        "party_address2": None,
        "party_zip": 700001,
        "party_gstin": "19AAACS1234A1Z5",
        "party_state": "WEST BENGAL",
        "consignee_name": "ANURASHI COMMOTRADE PVT. LTD",
        "branch_address1": "15, ALIPORE PARK ROAD, Kolkata",
        "branch_address2": None,
        "branch_zipcode": 700027,
        "branch_state": "WEST BENGAL",
        "godown_name": "Godown 2C",
        "purchase_ledger": "Purchase of Raw Jute",
        "claim_ledger": "Claim on Gross Purchase",
        "tds_ledger": "TDS ON PURCHASE OF GOODS (194Q)",
        "round_off_ledger": "Round Off",
    }
    h.update(overrides)
    return h


def _line(**overrides):
    """One MR line, shaped like get_jute_purchase_tally_line_query (kg, rate per quintal)."""
    ln = {
        "jute_mr_id": 1,
        "jute_mr_li_id": 10,
        "item_name": "TD5 (Desi)",
        "accepted_weight": 7262.0,
        "actual_qty": 50.0,
        "rate": 7424.0,
        "claim_rate": 30.0,
    }
    ln.update(overrides)
    return ln


def _sample_lines():
    return [
        _line(),
        _line(jute_mr_li_id=11, item_name="TD6 (Desi)", accepted_weight=2178.0,
              actual_qty=15.0, rate=7222.0, claim_rate=0.0),
    ]


def _dr_minus_cr(rows):
    return round(sum(
        r["G"] if r["H"] == "Dr" else -r["G"] for r in rows if r.get("F")
    ), 2)


class TestBuildPurchaseTallyRows:
    def test_sample_voucher_layout(self):
        rows = build_purchase_tally_rows([_header()], _sample_lines())

        # Party, Purchase + item 1, item 2, Claim, TDS, Round Off -- as in the sample sheet.
        assert [r.get("F") for r in rows] == [
            "Snehjyoti Commerce Pvt Ltd (J-P)",
            "Purchase of Raw Jute",
            None,
            "Claim on Gross Purchase",
            "TDS ON PURCHASE OF GOODS (194Q)",
            "Round Off",
        ]
        for r in rows:
            assert r["A"] == date(2026, 5, 23)            # supplier invoice date
            assert r["B"] == "Purchase"
            assert r["C"] == "ACPL/JMR/26-27/00024"        # full app MR number
            assert r["N"] == "Item Invoice"

        party, purchase, item2, claim, tds, ro = rows
        assert (purchase["G"], purchase["H"]) == (696426.04, "Dr")
        assert purchase["I"] == "TD5 (Desi)" and purchase["J"] == 7262.0
        assert purchase["K"] == 74.24 and purchase["L"] == "kg" and purchase["M"] == 539130.88
        assert purchase["O"] == "Godown 2C"
        assert purchase["AE"] == "Raw Jute: 50 Bales"
        assert purchase["AF"] == "ACPL/JMR/26-27/00024"
        assert item2["I"] == "TD6 (Desi)" and item2["M"] == 157295.16
        assert item2["AE"] == "Raw Jute: 15 Bales"
        assert (claim["G"], claim["H"]) == (2178.6, "Cr")
        assert claim["AG"] == "(Q): TD5 (Desi)- 7262 Kg. X Rs. 0.30/-"
        assert (tds["G"], tds["H"]) == (694.0, "Cr")
        # 696426.04 - 2178.6 - 694 = 693553.44 -> net 693553, Round Off 0.44 Cr.
        assert (ro["G"], ro["H"]) == (0.44, "Cr")
        assert (party["G"], party["H"]) == (693553.0, "Cr")
        assert _dr_minus_cr(rows) == 0

    def test_party_row_carries_supplier_consignee_and_narration(self):
        party = build_purchase_tally_rows([_header()], _sample_lines())[0]
        assert party["D"] == "7ram Kumar Rakhit Lane, Ground Floor, Kolkata"
        assert party["E"] == 700001 and party["U"] == 700001
        assert party["P"] == "SCP/MAY/021/26-27" and party["Q"] == date(2026, 5, 23)
        assert party["R"] == "Snehjyoti Commerce Pvt Ltd (J-P)"
        assert (party["S"], party["T"], party["V"], party["W"]) == (
            "India", "West Bengal", "Regular", "West Bengal")
        assert party["X"] == party["Y"] == "ANURASHI COMMOTRADE PVT. LTD"
        assert (party["Z"], party["AA"], party["AB"], party["AC"], party["AD"]) == (
            "Primary", "15, ALIPORE PARK ROAD, Kolkata", "West Bengal", "India", 700027)
        assert party["AH"] == (
            "Being 9440 Kg. Raw Jute Purchase From Snehjyoti Commerce Pvt Ltd (J-P) "
            "Against Inv: SCP/MAY/021/26-27,Dt: 23.05.2026,CH: MAY/021,Dt: 23.05.2026,"
            "MR.NO.ACPL/JMR/26-27/00024,Amt Rs. 693553/-"
        )
        assert not party.get("I")
        assert (party["AI"], party["AJ"]) == ("New Ref", "ACPL/JMR/26-27/00024")

    def test_bill_allocation_only_on_party_row(self):
        rows = build_purchase_tally_rows([_header()], _sample_lines())
        assert [r.get("AI") for r in rows] == ["New Ref"] + [None] * (len(rows) - 1)

    def test_plain_voucher_has_only_party_and_purchase(self):
        rows = build_purchase_tally_rows(
            [_header(tds_amount=0, unit_conversion="LOOSE")],
            [_line(accepted_weight=1000.0, rate=7500.0, claim_rate=0.0)],
        )
        assert [r["F"] for r in rows] == ["Snehjyoti Commerce Pvt Ltd (J-P)", "Purchase of Raw Jute"]
        assert rows[0]["G"] == rows[1]["G"] == 75000.0
        assert rows[1]["AE"] == "Raw Jute: Loose"

    def test_upward_round_off_is_debit(self):
        # 1001 kg x 74.99 = 75064.99 -> net 75065, Round Off 0.01 Dr.
        rows = build_purchase_tally_rows(
            [_header(tds_amount=0)],
            [_line(accepted_weight=1001.0, rate=7499.0, claim_rate=0.0)],
        )
        ro = rows[-1]
        assert (ro["F"], ro["G"], ro["H"]) == ("Round Off", 0.01, "Dr")
        assert rows[0]["G"] == 75065.0
        assert _dr_minus_cr(rows) == 0

    def test_invoice_date_falls_back_to_mr_date(self):
        rows = build_purchase_tally_rows(
            [_header(invoice_no=None, invoice_date=None, challan_no=None)], _sample_lines())
        assert rows[0]["A"] == date(2026, 6, 16)
        assert rows[0]["P"] is None and rows[0]["Q"] is None
        assert "Against Inv" not in rows[0]["AH"] and "CH:" not in rows[0]["AH"]

    def test_company_ledger_names_are_used(self):
        rows = build_purchase_tally_rows([_header(
            purchase_ledger="Raw Jute Purchase A/c", claim_ledger="Quality Claim",
            tds_ledger="TDS 194Q Payable", round_off_ledger="Rounding Off")], _sample_lines())
        assert [r.get("F") for r in rows][1:] == [
            "Raw Jute Purchase A/c", None, "Quality Claim", "TDS 194Q Payable", "Rounding Off"]

    def test_null_tds_is_derived_stored_tds_is_kept(self):
        calls = []

        def tds_for(h, gross):
            calls.append((h["jute_mr_id"], gross))
            return 696.43

        rows = build_purchase_tally_rows([_header(tds_amount=None)], _sample_lines(), tds_for)
        assert calls == [(1, 696426.04)]
        assert [r["G"] for r in rows if r.get("F") == "TDS ON PURCHASE OF GOODS (194Q)"] == [696.43]
        assert _dr_minus_cr(rows) == 0

        build_purchase_tally_rows([_header(tds_amount=0.0)], _sample_lines(), tds_for)
        assert len(calls) == 1                       # a stored 0 is not re-derived
        rows = build_purchase_tally_rows([_header(tds_amount=None)], _sample_lines())
        assert "TDS ON PURCHASE OF GOODS (194Q)" not in [r.get("F") for r in rows]

    def test_kg_rate_is_rounded_half_up_to_2dp_and_drives_amount(self):
        for rate_qtl, rate_kg in ((7499.5, 75.0), (7012.5, 70.13), (7100.5, 71.01), (7424.5, 74.25)):
            rows = build_purchase_tally_rows(
                [_header(tds_amount=0)], [_line(accepted_weight=1000.0, rate=rate_qtl, claim_rate=0.0)])
            assert rows[1]["K"] == rate_kg
            assert rows[1]["M"] == rows[0]["G"] == round(1000 * rate_kg, 2)

    def test_negative_amounts_flip_side_and_still_balance(self):
        # claim larger than the purchase -> net negative; a negative stored TDS.
        rows = build_purchase_tally_rows(
            [_header(tds_amount=-10.0)], [_line(accepted_weight=100.0, rate=100.0, claim_rate=500.0)])
        party, tds = rows[0], rows[-1]
        assert party["G"] > 0 and party["H"] == "Dr"
        assert (tds["F"], tds["G"], tds["H"]) == ("TDS ON PURCHASE OF GOODS (194Q)", 10.0, "Dr")
        assert all(r["G"] >= 0 for r in rows if r.get("F"))
        assert _dr_minus_cr(rows) == 0

    def test_no_gstin_leaves_registration_type_blank(self):
        party = build_purchase_tally_rows([_header(party_gstin=None)], _sample_lines())[0]
        assert party["V"] is None

    def test_legacy_quintal_mr_and_mr_without_lines_are_skipped(self):
        headers = [_header(jute_mr_id=1), _header(jute_mr_id=2), _header(jute_mr_id=3)]
        lines = [
            _line(jute_mr_id=1, accepted_weight=None),   # pre FY 26-27 import, quintals
            _line(jute_mr_id=3),
        ]
        rows = build_purchase_tally_rows(headers, lines)
        assert {r["C"] for r in rows} == {"ACPL/JMR/26-27/00024"}
        assert len([r for r in rows if r.get("F") == "Purchase of Raw Jute"]) == 1

    def test_zero_weight_line_is_dropped(self):
        rows = build_purchase_tally_rows(
            [_header(tds_amount=0)],
            [_line(claim_rate=0.0), _line(jute_mr_li_id=11, accepted_weight=0.0)],
        )
        assert len([r for r in rows if r.get("I")]) == 1


def _mock_result(rows):
    result = MagicMock()
    mocked = []
    for r in rows:
        m = MagicMock()
        m._mapping = r
        mocked.append(m)
    result.fetchall.return_value = mocked
    return result


class TestJuteTallyDownloadEndpoint:
    URL = "/api/juteReports/tally-download"

    @pytest.fixture(autouse=True)
    def setup_overrides(self):
        self.db = MagicMock()
        app.dependency_overrides[get_current_user_with_refresh] = lambda: {"user_id": 1}
        app.dependency_overrides[get_tenant_db] = lambda: self.db
        yield
        app.dependency_overrides.clear()

    @pytest.mark.parametrize("qs, field", [
        ("date_from=2026-01-01&date_to=2026-01-31", "co_id"),
        ("co_id=1&date_to=2026-01-31", "date_from"),
        ("co_id=1&date_from=2026-01-01", "date_to"),
        ("co_id=1&date_from=01-01-2026&date_to=2026-01-31", "date_from"),
        ("co_id=abc&date_from=2026-01-01&date_to=2026-01-31", "co_id"),
        ("co_id=1&branch_id=x&date_from=2026-01-01&date_to=2026-01-31", "branch_id"),
    ])
    def test_bad_params_return_400(self, qs, field):
        response = client.get(f"{self.URL}?{qs}")
        assert response.status_code == 400
        assert field in response.json()["detail"]

    def test_fills_the_template_workbook(self):
        self.db.execute.side_effect = [_mock_result([_header()]), _mock_result(_sample_lines())]
        response = client.get(f"{self.URL}?co_id=65&branch_id=30&date_from=2026-06-01&date_to=2026-06-30")
        assert response.status_code == 200
        assert "spreadsheetml" in response.headers["content-type"]

        params = self.db.execute.call_args_list[0].args[1]
        assert params == {"co_id": 65, "branch_id": 30,
                          "date_from": "2026-06-01", "date_to": "2026-06-30"}

        out = openpyxl.load_workbook(BytesIO(response.content))
        tpl = openpyxl.load_workbook(_PURCHASE_TALLY_TEMPLATE)
        assert out.sheetnames == tpl.sheetnames
        ws, tws = out["Accounting Voucher"], tpl["Accounting Voucher"]
        assert [c.value for c in ws[1]] == [c.value for c in tws[1]] + ["Bill Type of Ref", "Bill Name"]
        assert (ws["AI2"].value, ws["AJ2"].value) == ("New Ref", "ACPL/JMR/26-27/00024")
        assert ws["AI3"].value is None and ws["AJ3"].value is None
        assert ws.max_row == 7                       # header + 6 voucher rows, no sample rows left
        assert ws["C2"].value == "ACPL/JMR/26-27/00024"
        assert ws["F2"].value == "Snehjyoti Commerce Pvt Ltd (J-P)"
        assert ws["A2"].number_format == ws["Q2"].number_format == "d-mmm-yy"
        assert ws["A2"].value.date() == date(2026, 5, 23)
        assert ws["F7"].value == "Round Off"
        assert ws.cell(2, 9).value is None          # party row has no item
        rm, trm = out["Accounting Voucher (Read Me)"], tpl["Accounting Voucher (Read Me)"]
        assert [[c.value for c in r] for r in rm.iter_rows()] == \
               [[c.value for c in r] for r in trm.iter_rows()]

    def test_null_tds_is_derived_through_mr_logic(self, monkeypatch):
        calls = []

        def fake_derive(db, **kw):
            calls.append(kw)
            return {"tds_amount": 696.43}

        monkeypatch.setattr("src.juteProcurement.reports.derive_mr_tds", fake_derive)
        self.db.execute.side_effect = [_mock_result([_header(tds_amount=None)]), _mock_result(_sample_lines())]
        response = client.get(f"{self.URL}?co_id=65&date_from=2026-06-01&date_to=2026-06-30")
        assert response.status_code == 200, response.text
        assert calls == [{"party_id": "8668", "mr_date": date(2026, 6, 16), "mr_id": 1,
                          "branch_mr_no": 24, "total_amount": 696426.04}]
        ws = openpyxl.load_workbook(BytesIO(response.content))["Accounting Voucher"]
        assert ("TDS ON PURCHASE OF GOODS (194Q)", 696.43) in [
            (ws.cell(r, 6).value, ws.cell(r, 7).value) for r in range(2, ws.max_row + 1)]

    def test_empty_range_returns_headers_only(self):
        self.db.execute.side_effect = [_mock_result([]), _mock_result([])]
        response = client.get(f"{self.URL}?co_id=65&date_from=2026-06-01&date_to=2026-06-30")
        assert response.status_code == 200
        ws = openpyxl.load_workbook(BytesIO(response.content))["Accounting Voucher"]
        assert ws.max_row == 1


class TestPurchaseTallyQueries:
    def test_approved_only_and_tally_names(self):
        from src.juteProcurement.reportQueries import (
            get_jute_purchase_tally_header_query,
            get_jute_purchase_tally_line_query,
        )
        header = str(get_jute_purchase_tally_header_query())
        line = str(get_jute_purchase_tally_line_query())
        for sql in (header, line):
            assert "jm.status_id IN (3)" in sql
            assert "jm.jute_mr_date BETWEEN :date_from AND :date_to" in sql
        assert "'PARTY'" in header and "'GODOWN'" in header
        assert "'ITEM'" in line
        assert "mrli.status NOT IN ('4', '6')" in line
