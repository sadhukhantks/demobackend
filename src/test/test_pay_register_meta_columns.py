"""Tests for pay-register meta columns (PF/UAN/ESI), column layout and totals.

Covers the pure helpers in src.hrms.payRegister that drive how
tbl_payslip_print_component config becomes register columns:
  - META_FIELDS sentinel → source mapping
  - _meta_value            → text lookup for a meta column
  - _build_register_columns → ordered grid/export columns (NO per-row Total)
  - _register_cell_value    → one cell's value
  - _grand_total_row        → vertical grand-total row (numeric columns summed)
"""

from io import BytesIO

from openpyxl import load_workbook

from src.hrms.payRegister import (
    META_FIELDS,
    _meta_value,
    _build_register_columns,
    _register_cell_value,
    _grand_total_row,
    _register_workbook,
)


def _comp(cid, label, *, total_print=False, payslip_print=True, group=""):
    """Build a component dict shaped like _resolve_print_components output."""
    is_meta = cid < 0
    return {
        "component_id": cid,
        "label": label,
        "group": group,
        "total_print": total_print,
        "payslip_print": payslip_print,
        "is_meta": is_meta,
        "meta_key": META_FIELDS.get(cid),
    }


class TestMetaFieldMapping:
    def test_sentinels_map_to_expected_sources(self):
        assert META_FIELDS[-4] == "pf_no"
        assert META_FIELDS[-5] == "pf_uan_no"
        assert META_FIELDS[-6] == "esi_no"
        # identity fields sourced from the salary row
        assert META_FIELDS[-1] == "emp_code"
        assert META_FIELDS[-2] == "emp_name"
        assert META_FIELDS[-3] == "department_name"
        # the three that complete the legacy register's identity set
        assert META_FIELDS[-7] == "dept_code"
        assert META_FIELDS[-8] == "designation"
        assert META_FIELDS[-9] == "eb_id"


class TestMetaValue:
    def test_pf_uan_esi_read_from_meta_dict(self):
        emp = {"emp_code": "110199", "meta": {"pf_no": "PF1", "pf_uan_no": "UAN1", "esi_no": "ESI1"}}
        assert _meta_value(emp, "pf_no") == "PF1"
        assert _meta_value(emp, "pf_uan_no") == "UAN1"
        assert _meta_value(emp, "esi_no") == "ESI1"

    def test_identity_keys_read_from_employee_row(self):
        emp = {
            "emp_code": "110199", "emp_name": "Asha", "department_name": "HR",
            "dept_code": "069", "designation": "GARDENER", "meta": {},
        }
        assert _meta_value(emp, "emp_code") == "110199"
        assert _meta_value(emp, "department_name") == "HR"
        # dept_code / designation / eb_id ride on the salary row, not on emp['meta']
        assert _meta_value(emp, "dept_code") == "069"
        assert _meta_value(emp, "designation") == "GARDENER"
        assert _meta_value({**emp, "eb_id": 54539}, "eb_id") == "54539"

    def test_missing_meta_returns_empty_string_not_none(self):
        emp = {"emp_code": "1", "meta": {}}
        assert _meta_value(emp, "pf_no") == ""
        assert _meta_value(emp, None) == ""


class TestBuildRegisterColumns:
    def test_configured_order_with_anchors_meta_no_per_row_total(self):
        components = [
            _comp(-1, "Emp Code"),          # anchor — must be skipped (already leading)
            _comp(101, "Basic", total_print=True),
            _comp(-4, "PF No"),
            _comp(-6, "ESI No"),
            _comp(102, "Net", total_print=True),
        ]
        cols = _build_register_columns(components, used_config=True)
        keys = [c["key"] for c in cols]
        # Emp Code + Name always lead; configured -1 is not duplicated;
        # there is NO per-row __total__ column.
        assert keys[:2] == ["emp_code", "emp_name"]
        assert keys == ["emp_code", "emp_name", "c101", "m4", "m6", "c102"]
        assert "__total__" not in keys
        by_key = {c["key"]: c for c in cols}
        assert by_key["m4"]["kind"] == "text"
        assert by_key["m4"]["meta_key"] == "pf_no"
        assert by_key["c101"]["kind"] == "num"

    def test_fallback_includes_department_and_all_components(self):
        components = [_comp(101, "Basic"), _comp(102, "HRA")]
        cols = _build_register_columns(components, used_config=False)
        keys = [c["key"] for c in cols]
        assert keys == ["emp_code", "emp_name", "department_name", "c101", "c102"]
        assert "__total__" not in keys

    def test_configured_anchor_labels_and_repeated_component(self):
        # The legacy contractor sheet heads the anchors 'eb_no' / 'emp_name' and
        # prints component 66 twice (Amount, Gross1): labels follow the config,
        # the second key gets a suffix so the grid's fields stay distinct.
        components = [
            _comp(-1, "eb_no"), _comp(-2, "emp_name"),
            _comp(66, "Amount"), _comp(7, "Basic"), _comp(66, "Gross1"),
        ]
        cols = _build_register_columns(components, used_config=True)
        assert [c["label"] for c in cols] == ["eb_no", "emp_name", "Amount", "Basic", "Gross1"]
        assert [c["key"] for c in cols] == ["emp_code", "emp_name", "c66", "c7", "c66_2"]
        assert cols[-1]["component_id"] == 66
        # Without -1 / -2 rows the anchors keep their default headings.
        plain = _build_register_columns([_comp(7, "Basic")], used_config=True)
        assert [c["label"] for c in plain][:2] == ["Emp Code", "Name"]


class TestRegisterCellValue:
    def test_meta_text_and_component_number(self):
        emp = {
            "emp_code": "110199",
            "emp_name": "Asha",
            "values": {101: 5000.0},
            "meta": {"pf_no": "PF1"},
        }
        pf_col = {"key": "m4", "kind": "text", "meta_key": "pf_no"}
        basic_col = {"key": "c101", "kind": "num", "component_id": 101}
        anchor_col = {"key": "emp_code", "kind": "text"}

        assert _register_cell_value(pf_col, emp) == "PF1"
        assert _register_cell_value(basic_col, emp) == 5000.0
        assert _register_cell_value(anchor_col, emp) == "110199"

    def test_missing_component_value_is_zero(self):
        emp = {"values": {}, "meta": {}}
        col = {"key": "c999", "kind": "num", "component_id": 999}
        assert _register_cell_value(col, emp) == 0.0


class TestGrandTotalRow:
    def test_sums_numeric_columns_and_labels_name(self):
        components = [
            _comp(101, "Basic"),
            _comp(-4, "PF No"),
            _comp(102, "HRA"),
        ]
        cols = _build_register_columns(components, used_config=True)
        employees = [
            {"emp_code": "1", "emp_name": "A", "values": {101: 5000.0, 102: 1000.0}, "meta": {"pf_no": "PF1"}},
            {"emp_code": "2", "emp_name": "B", "values": {101: 3000.0, 102: 500.0}, "meta": {"pf_no": "PF2"}},
        ]
        row = _grand_total_row(cols, employees)
        assert row["emp_name"] == "GRAND TOTAL"
        assert row["emp_code"] == ""        # leading text column blank
        assert row["m4"] == ""              # meta text column never summed
        assert row["c101"] == 8000.0
        assert row["c102"] == 1500.0

    def test_empty_numeric_columns_sum_to_zero(self):
        components = [_comp(101, "Basic")]
        cols = _build_register_columns(components, used_config=True)
        employees = [{"emp_code": "1", "emp_name": "A", "values": {}, "meta": {}}]
        row = _grand_total_row(cols, employees)
        assert row["c101"] == 0.0


class TestRegisterWorkbook:
    """The .xlsx follows the legacy sheet: company / title / headers / data /
    grand total, with the legacy identity columns configured as meta rows."""

    def _ctx(self):
        # The contractor register's legacy identity set, in legacy order, then
        # two components — the shape the sls seed migration writes.
        components = [
            _comp(-1, "Eb No"), _comp(-2, "Worker Name"), _comp(-7, "Dept Code"),
            _comp(-3, "Department"), _comp(-8, "Designation"),
            _comp(-5, "UAN No"), _comp(-6, "ESI No"), _comp(-9, "Eb Id"),
            _comp(70, "Rate", total_print=True), _comp(21, "Net", total_print=True),
        ]
        employees = [
            {"eb_id": 1, "emp_code": "G980015", "emp_name": "ANIL", "department_name": "GARDNERS",
             "dept_code": "069", "designation": "GARDENER",
             "values": {70: 500.0, 21: 12000.0}, "meta": {"pf_uan_no": "UAN1", "esi_no": "ESI1"}},
            {"eb_id": 2, "emp_code": "G980016", "emp_name": "BINOD", "department_name": "GARDNERS",
             "dept_code": "069", "designation": "GARDENER",
             "values": {70: 500.0, 21: 8000.5}, "meta": {}},
        ]
        return {
            "company_name": "NJM",
            "register_title": "NJM_CONT_WAGES_NEW PAY REGISTER FOR THE PERIOD FROM 01-08-2026 TO 31-08-2026",
            "used_config": True, "components": components, "employees": employees,
        }

    def test_legacy_layout(self):
        ws = load_workbook(BytesIO(_register_workbook(self._ctx()))).active
        # openpyxl reads a written "" back as None; normalise so blanks compare.
        rows = [["" if c.value is None else c.value for c in r] for r in ws.iter_rows()]
        assert ws.title == "Pay Register"
        assert rows[0][0] == "NJM"
        assert rows[1][0] == "NJM_CONT_WAGES_NEW PAY REGISTER FOR THE PERIOD FROM 01-08-2026 TO 31-08-2026"
        # Row 3 headers: anchors (headed by their configured labels), then the
        # configured columns in order; no blank spacer row like the old layout.
        assert rows[2] == ["Eb No", "Worker Name", "Dept Code", "Department", "Designation",
                           "UAN No", "ESI No", "Eb Id", "Rate", "Net"]
        assert rows[3] == ["G980015", "ANIL", "069", "GARDNERS", "GARDENER", "UAN1", "ESI1", "1", 500.0, 12000.0]
        assert rows[4] == ["G980016", "BINOD", "069", "GARDNERS", "GARDENER", "", "", "2", 500.0, 8000.5]
        # Grand total closes the sheet: numeric columns summed, text blank.
        assert rows[5] == ["", "GRAND TOTAL", "", "", "", "", "", "", 1000.0, 20000.5]
        assert len(rows) == 6

    def test_no_employees_still_writes_headers(self):
        ctx = self._ctx()
        ctx["employees"] = []
        ws = load_workbook(BytesIO(_register_workbook(ctx))).active
        rows = [[c.value for c in r] for r in ws.iter_rows()]
        assert rows[2][:3] == ["Eb No", "Worker Name", "Dept Code"]
        assert len(rows) == 3
