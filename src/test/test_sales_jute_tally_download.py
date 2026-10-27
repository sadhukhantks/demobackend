"""
Tests for the Sales Raw Jute Tally export.
src/sales/reports.py — build_tally_voucher_rows + GET /api/salesReports/jute-tally-download
"""

from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
from unittest.mock import MagicMock

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.authorization.utils import get_current_user_with_refresh
from src.config.db import get_tenant_db
from src.main import app
from src.sales.reportQueries import (
    get_sales_jute_tally_header_query,
    get_sales_jute_tally_line_query,
)
from src.sales.reports import build_tally_voucher_rows

client = TestClient(app)

URL = "/api/salesReports/jute-tally-download"

# docs/Sales.xlsx "Accounting Voucher" header row, VERBATIM (lowercase
# "change mode", typos at AG / AJ) minus the dropped duplicates "Change Mode ",
# "New Ref No." and "Delivery Note Date/...". Duplicated so any drift fails.
_EXPECTED_HEADERS = [
    "Voucher Date", "Voucher Type Name", "Voucher Number",
    "Buyer/Supplier - Address", "Buyer/Supplier - Pincode", "Ledger Name",
    "Ledger Amount", "Ledger Amount Dr/Cr", "Item Name", "item description",
    "Billed Quantity", "Item Rate", "Item Rate per", "Item Amount",
    "Buyer/Supplier - Bill to/from", "Dispatched through",
    "Destination", "Reference No.", "batch lot no", "Godown",
    "Despatch Doc No", "Consignee (ship to)",
    "Description of Ledger", "change mode", "order no", "order dt",
    "terms of payment", "terms of delivery", "state", "ship to adrress",
    "Voucher Narration", "Place of supply", "state consginee", "consignee zip",
]

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _header(**over):
    h = {
        "invoice_id": 1,
        "invoice_date": date(2026, 10, 5),
        "voucher_no": "ACP/OCT/001/25-26",
        "vehicle_no": "AP 05TF 0259",
        "challan_no": " OCT/001 ",
        "challan_date": date(2026, 10, 5),
        "buyer_order_no": "123",
        "buyer_order_date": date(2025, 6, 10),
        "payment_terms": 90,
        "round_off": 0,
        "party_name": "Jagrati Trade Services Pvt. Ltd.",
        "consignee_name": "Jagrati Trade Services Pvt. Ltd.",
        "godown_name": "Godown 1C",
        "sales_ledger": "Sale of Raw Jute",
        "claim_ledger": "Claim on Gross Sales",
        "round_off_ledger": "Round Off",
        "bill_address": "3a shakespeare sarani, 8th floor",
        "bill_zip": 700071,
        "bill_state": "WEST BENGAL",
        "ship_address": "3q shakespeare sarani, 8 th floor",
        "ship_zip": 700071,
        "ship_state": "WEST BENGAL",
        "dispatch_address": "7 Ram Rakhit Lane, Kolkata",
        "destination": "3q shakespeare sarani, 8 th floor",
        "unit_conversion": "LOOSE",
        "lot_no": "ACPL/DVF/JMR/26-27/00025",
    }
    h.update(over)
    return h


def _line(**over):
    ln = {
        "invoice_id": 1,
        "invoice_line_item_id": 1,
        "item_name": "TD8 (A)",
        "quantity": 13950.0,
        "rate": 100.1,
        "amount_without_tax": 1396395.0,
        "uom_id": 163,
        "uom_name": None,  # 163 is missing from sls uom_mst
        "claim_amount": 4185.0,
        "bales": None,
    }
    ln.update(over)
    return ln


def _sample_lines():
    """The two lines behind the docs/Sales.xlsx sample voucher."""
    return [
        _line(),
        _line(invoice_line_item_id=2, item_name="TD9 (A)", quantity=10123.0,
              rate=98.0, amount_without_tax=992054.0, claim_amount=3036.9),
    ]


def _balance(rows):
    sums = defaultdict(float)
    for r in rows:
        if r.get("G") is not None:
            sums[r["C"]] += r["G"]
    return {k: round(v, 2) for k, v in sums.items()}


# ---------------------------------------------------------------------------
# Builder
# ---------------------------------------------------------------------------

class TestBuildTallyVoucherRows:

    def test_reproduces_sample_voucher(self):
        rows = build_tally_voucher_rows([_header()], _sample_lines())

        assert [r.get("F") or r.get("I") for r in rows] == [
            "Jagrati Trade Services Pvt. Ltd.", "Sale of Raw Jute",
            "TD8 (A)", "TD9 (A)", "Claim on Gross Sales",
        ]
        party, sale, item1, item2, claim = rows
        assert all(r["B"] == "Sales" and r["C"] == "ACP/OCT/001/25-26"
                   and r["AA"] == "Item Invoice" and r["A"] == date(2026, 10, 5)
                   for r in rows)

        assert party["G"] == -2381227.1
        assert sale["G"] == 2388449
        assert (item1["K"], item1["L"], item1["N"]) == (13950, 100.1, 1396395)
        assert (item2["K"], item2["L"], item2["N"]) == (10123, 98, 992054)
        assert claim["G"] == -7221.9
        assert claim["Z"] == "(Q): 24073 Kg. X Rs.0.30/-"
        assert party["AH"] == (
            "Being 24073 Kg. Raw Jute Sold To 'Jagrati Trade Services Pvt. Ltd.' "
            "against Inv No: ACP/OCT/001/25-26, Dt: 05.10.2026, "
            "REF NO: ACPL/DVF/JMR/26-27/00025, Amount Rs. 2381227.10/-"
        )

        # party-row header fields
        assert party["Q"] == "Truck No: AP 05TF 0259"
        assert "U" not in party  # New Ref No. column removed
        assert party["W"] == "Ch.No: OCT/001 , Dt: 05.10.2026"
        assert "Y" not in party  # Delivery Note Date column removed (dup of W)
        assert party["AC"] == "10.06.2025"
        assert party["AD"] == "90 days"
        assert party["AF"] == party["AI"] == party["AJ"] == "West Bengal"
        assert party["E"] == party["AK"] == 700071

        # item-row fields
        assert item1["J"] == "Raw Jute: Loose"
        assert item1["M"] == "kg"
        assert item1["S"] == "REF NO: ACPL/DVF/JMR/26-27/00025"
        assert item1["T"] == "ACPL/DVF/JMR/26-27/00025"
        assert item1["V"] == "Godown 1C"

        assert _balance(rows) == {"ACP/OCT/001/25-26": 0}

    def test_company_ledger_names_are_used(self):
        rows = build_tally_voucher_rows(
            [_header(round_off=-0.24, sales_ledger="Raw Jute Sales A/c",
                     claim_ledger="Quality Claim (Sales)", round_off_ledger="Rounding Off")],
            [_line(claim_amount=10.0)],
        )
        names = [r.get("F") for r in rows if r.get("F")]
        assert names[1:] == ["Raw Jute Sales A/c", "Quality Claim (Sales)", "Rounding Off"]

    def test_header_query_joins_company_ledgers(self):
        sql = str(get_sales_jute_tally_header_query())
        for t in ("'SALES_LEDGER'", "'SALES_CLAIM'", "'ROUND_OFF'"):
            assert t in sql
        assert "'Sale of Raw Jute'" in sql

    def test_lot_no_is_purchase_sheet_mr_number(self):
        # Batch/Lot No. must equal the purchase Tally sheet's MR no (AF), which
        # is the formatted MR number keyed on the MR's own branch + company.
        sql = get_sales_jute_tally_header_query().text
        assert "'JMR/'" in sql
        assert "LPAD(jm.branch_mr_no, 5, '0')" in sql
        assert "COALESCE(mcm.co_prefix, '')" in sql
        assert "COALESCE(mbm.branch_prefix, '')" in sql
        assert "LEFT JOIN co_mst mcm ON mcm.co_id = mbm.co_id" in sql
        assert "CONCAT(COALESCE(mbm.branch_prefix, ''), jm.branch_mr_no" not in sql

    def test_lot_no_prefers_receiving_company_mr(self):
        # The buyer's MR (juteTransfer chain hop) wins; the seller's own MR
        # (sales_invoice_jute.mr_id) is only the fallback.
        sql = get_sales_jute_tally_header_query().text
        assert "LEFT JOIN jute_mr rm ON rm.jute_mr_id = (" in sql
        assert "CASE WHEN rm.jute_mr_id IS NOT NULL THEN" in sql
        assert "LPAD(rm.branch_mr_no, 5, '0')" in sql
        assert "COALESCE(rcm.co_prefix, '')" in sql and "COALESCE(rbm.branch_prefix, '')" in sql
        # chain membership + the three match rules, in priority order
        assert "h.src_jute_mr_id = COALESCE(jm.src_jute_mr_id, jm.jute_mr_id)" in sql
        assert "h.jute_mr_id <> jm.jute_mr_id" in sql
        assert "SUBSTRING_INDEX(h.invoice_no, '/SI/', -1)" in sql
        assert "h.src_com_id = bm.co_id" in sql
        assert sql.index("ORDER BY (h.invoice_no =") < sql.index("ORDER BY ih.invoice_date")
        assert "LIMIT 1" in sql
        assert "WHEN jm.jute_mr_id IS NOT NULL THEN" in sql  # fallback

    def test_destination_column_is_the_resolved_destination(self):
        rows = build_tally_voucher_rows([_header(destination="15/16 B.T. Road")], [_line()])
        assert rows[0]["R"] == "15/16 B.T. Road"
        assert rows[0]["AG"] == "3q shakespeare sarani, 8 th floor"  # ship-to address untouched
        assert build_tally_voucher_rows([_header(destination=None)], [_line()])[0]["R"] == ""

    def test_header_query_destination_precedence(self):
        # buyer override > receiving mill (chain root branch) > ship-to address
        sql = get_sales_jute_tally_header_query().text
        assert "COALESCE(tldp.tally_name, tld.tally_name, spb.address, '') AS destination" in sql
        assert "LEFT JOIN jute_mr rt ON rt.jute_mr_id = COALESCE(jm.src_jute_mr_id, jm.jute_mr_id)" in sql
        assert "tld.co_id = rtbm.co_id" in sql
        assert "tld.link_type = 'DESTINATION'" in sql
        assert "tld.ref_id = rt.branch_id" in sql
        assert "tldp.co_id = bm.co_id" in sql
        assert "tldp.link_type = 'DESTINATION_PARTY'" in sql
        assert "tldp.ref_id = ih.party_id" in sql

    def test_dispatched_through_blank_when_goods_did_not_move(self):
        moved = build_tally_voucher_rows([_header()], [_line()])
        assert moved[0]["Q"] == "Truck No: AP 05TF 0259"
        same = build_tally_voucher_rows(
            [_header(dispatch_address="3Q SHAKESPEARE SARANI, 8TH FLOOR.",
                     ship_address="3q shakespeare sarani, 8 th floor")],
            [_line()],
        )
        assert same[0]["Q"] == ""
        blank_both = build_tally_voucher_rows(
            [_header(dispatch_address=None, ship_address=None)], [_line()]
        )
        assert blank_both[0]["Q"] == "Truck No: AP 05TF 0259"

    def test_claim_and_round_off_balance(self):
        # sls invoice 18049 shape: gross juteTransfer row, claim on one line.
        rows = build_tally_voucher_rows(
            [_header(round_off=-0.24, unit_conversion="BALE")],
            [
                _line(quantity=4785.0, rate=123.42, amount_without_tax=590564.7,
                      claim_amount=0.0, bales=32),
                _line(invoice_line_item_id=2, quantity=5383.0, rate=121.38,
                      amount_without_tax=653388.54, claim_amount=2691.5, bales=36),
            ],
        )
        assert [r.get("F") or r.get("I") for r in rows][-2:] == [
            "Claim on Gross Sales", "Round Off",
        ]
        assert rows[0]["G"] == -1241261.5
        assert rows[1]["G"] == 1243953.24
        assert rows[-1]["G"] == -0.24
        assert rows[2]["J"] == "Raw Jute: 32 Bales"
        assert rows[-2]["Z"] == "(Q): 5383 Kg. X Rs.0.50/-"
        assert _balance(rows) == {"ACP/OCT/001/25-26": 0}

    def test_no_claim_no_round_off_rows_when_zero(self):
        lines = [_line(claim_amount=0.0)]
        rows = build_tally_voucher_rows([_header()], lines)
        assert [r.get("F") or r.get("I") for r in rows] == [
            "Jagrati Trade Services Pvt. Ltd.", "Sale of Raw Jute", "TD8 (A)",
        ]

    def test_zero_qty_line_skipped(self):
        lines = _sample_lines() + [
            _line(invoice_line_item_id=3, item_name="ZERO", quantity=0.0,
                  rate=0.0, amount_without_tax=0.0, claim_amount=0.0),
        ]
        rows = build_tally_voucher_rows([_header()], lines)
        assert "ZERO" not in [r.get("I") for r in rows]
        assert _balance(rows) == {"ACP/OCT/001/25-26": 0}

    def test_quintal_lines_converted_to_kg(self):
        rows = build_tally_voucher_rows([_header(lot_no="")], [
            _line(uom_id=130, uom_name="QNT", quantity=75.5, rate=9899.89,
                  amount_without_tax=747441.7, claim_amount=0.0),
            _line(invoice_line_item_id=2, uom_id=None, uom_name=None,
                  quantity=0.51, rate=5000.0, amount_without_tax=2550.0,
                  claim_amount=0.0),
        ])
        items = [r for r in rows if r.get("I")]
        assert (items[0]["K"], items[0]["L"]) == (7550, 98.9989)
        assert (items[1]["K"], items[1]["L"]) == (51, 50)
        assert rows[0]["AH"].startswith("Being 7601 Kg.")
        # no lot -> no REF NO anywhere
        assert "REF NO" not in rows[0]["AH"]
        assert items[0]["S"] == items[0]["T"] == ""

    def test_invoice_without_lines_skipped(self):
        rows = build_tally_voucher_rows(
            [_header(invoice_id=9, round_off=0.12), _header()], _sample_lines()
        )
        assert {r["C"] for r in rows} == {"ACP/OCT/001/25-26"}
        assert len(rows) == 5

    def test_blank_optional_header_fields(self):
        rows = build_tally_voucher_rows(
            [_header(vehicle_no=None, challan_no=None, payment_terms=None,
                     buyer_order_date=None, godown_name="")],
            [_line(claim_amount=0.0)],
        )
        party, _, item = rows
        assert party["Q"] == party["W"] == party["AD"] == party["AC"] == ""
        assert item["V"] == ""
        dateless = build_tally_voucher_rows([_header(challan_date=None)], [_line()])
        assert dateless[0]["W"] == "Ch.No: OCT/001"

    def test_tally_names_fall_back_to_erp_names(self):
        header_sql = get_sales_jute_tally_header_query().text
        line_sql = get_sales_jute_tally_line_query().text
        assert "COALESCE(tlp.tally_name, pm.supp_name, '')" in header_sql
        assert "COALESCE(tlc.tally_name, cpm.supp_name, '')" in header_sql
        assert "COALESCE(tlg.tally_name, '')" in header_sql
        assert "COALESCE(tli.tally_name, im.item_name, '')" in line_sql
        for alias, sql in (("tlp", header_sql), ("tlg", header_sql), ("tli", line_sql)):
            assert f"{alias}.active = 1" in sql
            assert f"{alias}.co_id = bm.co_id" in sql


# ---------------------------------------------------------------------------
# Endpoint
# ---------------------------------------------------------------------------

def _result(dicts):
    res = MagicMock()
    rows = []
    for d in dicts:
        r = MagicMock()
        r._mapping = d
        rows.append(r)
    res.fetchall.return_value = rows
    return res


class TestSalesJuteTallyDownload:

    @pytest.fixture(autouse=True)
    def setup_overrides(self):
        self._mock_session = MagicMock()
        app.dependency_overrides[get_current_user_with_refresh] = lambda: {"user_id": 1}
        app.dependency_overrides[get_tenant_db] = lambda: self._mock_session
        yield
        app.dependency_overrides.clear()

    def _serve(self, headers, lines):
        self._mock_session.execute.side_effect = [_result(headers), _result(lines)]

    def test_missing_co_id_returns_400(self):
        r = client.get(f"{URL}?date_from=2026-01-01&date_to=2026-01-31")
        assert r.status_code == 400
        assert "co_id" in r.json()["detail"].lower()

    def test_missing_date_from_returns_400(self):
        r = client.get(f"{URL}?co_id=1&date_to=2026-01-31")
        assert r.status_code == 400
        assert "date_from" in r.json()["detail"].lower()

    def test_bad_date_format_returns_400(self):
        r = client.get(f"{URL}?co_id=1&date_from=01-01-2026&date_to=2026-01-31")
        assert r.status_code == 400
        assert "date_from" in r.json()["detail"].lower()

    def test_bad_branch_id_returns_400(self):
        r = client.get(f"{URL}?co_id=1&branch_id=x&date_from=2026-01-01&date_to=2026-01-31")
        assert r.status_code == 400

    def test_happy_path_xlsx(self):
        self._serve([_header()], _sample_lines())
        r = client.get(f"{URL}?co_id=1&branch_id=1&date_from=2026-10-01&date_to=2026-10-31")

        assert r.status_code == 200
        assert r.headers["content-type"].startswith(_XLSX_MIME)
        assert "attachment" in r.headers["content-disposition"].lower()

        wb = load_workbook(BytesIO(r.content))
        assert wb.sheetnames == ["Accounting Voucher"]
        ws = wb["Accounting Voucher"]
        assert [c.value for c in ws[1]] == _EXPECTED_HEADERS
        assert ws.max_row == 6  # header + party, sale, 2 items, claim

        assert isinstance(ws["A2"].value, datetime)
        assert ws["A2"].value.date() == date(2026, 10, 5)
        assert ws["A2"].number_format == "dd-mm-yyyy"
        assert ws["G2"].value == -2381227.1
        assert ws["G2"].number_format == "0.00"
        assert ws["H2"].value is None
        assert ws["K4"].value == 13950 and ws["K4"].number_format == "#,##0"
        assert ws["L4"].number_format == "0.00"
        assert ws["F6"].value == "Claim on Gross Sales"
        assert sum(ws[f"G{i}"].value or 0 for i in range(2, 7)) == pytest.approx(0, abs=0.001)

    def test_empty_range_returns_header_only(self):
        self._serve([], [])
        r = client.get(f"{URL}?co_id=1&date_from=2026-01-01&date_to=2026-01-31")
        assert r.status_code == 200
        ws = load_workbook(BytesIO(r.content))["Accounting Voucher"]
        assert ws.max_row == 1
        assert [c.value for c in ws[1]] == _EXPECTED_HEADERS

    def test_over_cap_returns_400(self):
        # 3 output rows per invoice (party, sale, item) -> 10,002 rows.
        n = 3334
        self._serve(
            [_header(invoice_id=i, voucher_no=f"V{i}") for i in range(n)],
            [_line(invoice_id=i, claim_amount=0.0) for i in range(n)],
        )
        r = client.get(f"{URL}?co_id=1&date_from=2026-01-01&date_to=2026-01-31")
        assert r.status_code == 400
        assert "too large" in r.json()["detail"].lower()
