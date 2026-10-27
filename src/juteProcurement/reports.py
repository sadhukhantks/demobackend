"""
Jute Procurement Report endpoints.
Provides endpoints for jute stock and related reports.
"""

import logging
from copy import copy
from io import BytesIO
from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from src.config.db import get_tenant_db
from src.authorization.utils import get_current_user_with_refresh
from src.juteProcurement.reportQueries import (
    get_jute_stock_report_query,
    get_jute_qty_wise_report_query,
    get_jute_txn_summary_query,
    get_jute_period_wise_query,
    get_jute_mr_in_stock_query,
    get_jute_mr_wise_query,
    get_jute_godown_wise_query,
    get_jute_with_value_query,
    get_jute_percent_claims_query,
    get_jute_mukham_moisture_query,
    get_batch_cost_report_query,
    get_batch_cost_plan_coverage_query,
    get_jute_issue_yarn_wise_query,
    get_mr_list_query,
    get_mr_list_count_query,
    get_mr_list_statuses_query,
    get_jute_purchase_tally_header_query,
    get_jute_purchase_tally_line_query,
    get_jute_mr_options_query,
    get_jute_mr_issue_detail_query,
    get_jute_mr_movement_detail_query,
    get_jute_po_options_query,
    get_jute_po_mr_receipt_detail_query,
    get_jute_po_pending_query,
)
from src.juteProcurement.batchPlanCoverage import summarize_plan_coverage
from src.juteProcurement.issuePlanner import (
    UnknownBranchError,
    build_issue_planner_report,
)
from collections import defaultdict
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import load_workbook

from src.juteProcurement.mr import derive_mr_tds
from src.juteProcurement.totals import compute_jute_totals
from src.sales.reports import _dmy, _kg_text

logger = logging.getLogger(__name__)

router = APIRouter()


@router.get("/stock")
def get_jute_stock_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """
    Get daily jute stock position report for a given branch and date.

    Returns opening stock, receipt, issue, closing stock, and MTD receipt/issue
    grouped by item group and item.

    Query params:
    - branch_id: Branch ID (required)
    - date: Report date in YYYY-MM-DD format (required)
    - basis: "actual" (default) or "accounts" (decision D5, 2026-09-03)
    """
    try:
        q_branch_id = request.query_params.get("branch_id")
        q_date = request.query_params.get("date")

        if not q_branch_id:
            raise HTTPException(status_code=400, detail="branch_id is required")
        if not q_date:
            raise HTTPException(status_code=400, detail="date is required")

        try:
            branch_id = int(q_branch_id)
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid branch_id")

        q_basis = (request.query_params.get("basis") or "actual").lower()
        if q_basis not in ("actual", "accounts"):
            raise HTTPException(status_code=400, detail="basis must be 'actual' or 'accounts'")

        query = get_jute_stock_report_query(basis=q_basis)
        rows = db.execute(query, {
            "branch_id": branch_id,
            "report_date": _parse_iso_date(q_date, "date"),
        }).fetchall()

        data = [dict(r._mapping) for r in rows]

        return {"data": data}

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute stock report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute stock report",
        )


@router.get("/batch-cost")
def get_batch_cost_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """
    Get batch cost report: yarn quality-wise planned vs actual jute issue.

    Compares planned issue (from batch daily assignments + batch plan percentages)
    against actual issue for a given branch and date. Returns planned weight,
    actual weight, average rate, issue value, and variance per quality — for the
    date itself and month-to-date (1st of that month through the date,
    TKT-000213) as the mtd_* columns.

    Query params:
    - branch_id: Branch ID (required)
    - date: Report date in YYYY-MM-DD format (required)
    - basis: "actual" (default) or "accounts" (decision D5, 2026-09-03)

    Returns {"data": [...], "month_start": "YYYY-MM-01", "report_date": "YYYY-MM-DD",
    "plan_status": {...}} — plan_status.messages explains a 0 planned weight
    (no batch plan assigned / nothing issued, TKT-000206).
    """
    try:
        q_branch_id = request.query_params.get("branch_id")
        q_date = request.query_params.get("date")

        if not q_branch_id:
            raise HTTPException(status_code=400, detail="branch_id is required")
        if not q_date:
            raise HTTPException(status_code=400, detail="date is required")

        try:
            branch_id = int(q_branch_id)
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid branch_id")

        q_basis = (request.query_params.get("basis") or "actual").lower()
        if q_basis not in ("actual", "accounts"):
            raise HTTPException(status_code=400, detail="basis must be 'actual' or 'accounts'")

        report_date = datetime.strptime(
            _parse_iso_date(q_date, "date"), "%Y-%m-%d"
        ).date()
        month_start = report_date.replace(day=1)

        params = {
            "branch_id": branch_id,
            "month_start": month_start.isoformat(),
            "report_date": report_date.isoformat(),
        }
        rows = db.execute(get_batch_cost_report_query(basis=q_basis), params).fetchall()
        data = [dict(r._mapping) for r in rows]

        # TKT-000206: say why planned weight is 0 (no batch plan assigned for
        # the day / a yarn quality, or nothing issued).
        coverage = db.execute(
            get_batch_cost_plan_coverage_query(basis=q_basis), params
        ).fetchall()
        plan_status = summarize_plan_coverage(
            [dict(r._mapping) for r in coverage], report_date.isoformat()
        )

        return {
            "data": data,
            "month_start": month_start.isoformat(),
            "report_date": report_date.isoformat(),
            "plan_status": plan_status,
        }

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching batch cost report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching batch cost report",
        )


@router.get("/issue-yarn-wise")
def get_jute_issue_yarn_wise_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """
    Jute Issue — Yarn Quality Wise (TKT-000214): for each yarn quality, the
    jute issued on the report date and month-to-date (1st of that month up to
    and including the date), in kg. Each figure is raw jute + wastage and comes
    with its wastage part, so the page can exclude wastage on the client.

    Query params:
    - branch_id: Branch ID (required)
    - date: Report date in YYYY-MM-DD format (required)

    Returns {"data": [{yarn_type_id, yarn_type_name, day_issue_kg,
    mtd_issue_kg, day_waste_kg, mtd_waste_kg}], "month_start": "YYYY-MM-01",
    "report_date": "YYYY-MM-DD"}.
    """
    try:
        branch_id = _require_branch_id(request)
        q_date = request.query_params.get("date")
        if not q_date:
            raise HTTPException(status_code=400, detail="date is required")
        try:
            report_date = datetime.strptime(q_date, "%Y-%m-%d").date()
        except (TypeError, ValueError):
            raise HTTPException(
                status_code=400, detail="Invalid date format, expected YYYY-MM-DD"
            )
        month_start = report_date.replace(day=1)

        rows = db.execute(get_jute_issue_yarn_wise_query(), {
            "branch_id": branch_id,
            "month_start": month_start.isoformat(),
            "report_date": report_date.isoformat(),
        }).fetchall()

        return {
            "data": [dict(r._mapping) for r in rows],
            "month_start": month_start.isoformat(),
            "report_date": report_date.isoformat(),
        }

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute issue yarn wise report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute issue yarn wise report",
        )


@router.get("/issue-planner")
def get_jute_issue_planner_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """
    Jute Issue Planner: estimated godown-to-batching issue for the report date.

    Combines previous-day spinning production, the batch plan window
    (report date .. +4 days), the 06:00 spreader roll-stock snapshot, and
    maturity hours into a per-quality required issue (MT and bales), plus
    godown stock coverage in days.

    Query params:
    - branch_id: Branch ID (required)
    - date: Issue date in YYYY-MM-DD format (required)
    """
    try:
        q_branch_id = request.query_params.get("branch_id")
        q_date = request.query_params.get("date")

        if not q_branch_id:
            raise HTTPException(status_code=400, detail="branch_id is required")
        if not q_date:
            raise HTTPException(status_code=400, detail="date is required")

        try:
            branch_id = int(q_branch_id)
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid branch_id")

        try:
            report_date = datetime.strptime(q_date, "%Y-%m-%d").date()
        except ValueError:
            raise HTTPException(
                status_code=400, detail="Invalid date format, expected YYYY-MM-DD"
            )

        try:
            report = build_issue_planner_report(db, branch_id, report_date)
        except UnknownBranchError as e:
            raise HTTPException(status_code=404, detail=str(e))

        return {"data": report}

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute issue planner report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail=f"Error fetching jute issue planner report: {str(e)}",
        )


def _parse_iso_date(value: str, field_name: str) -> str:
    """Validate a YYYY-MM-DD date string. Raises HTTPException(400) if invalid."""
    try:
        datetime.strptime(value, "%Y-%m-%d")
    except (TypeError, ValueError):
        raise HTTPException(
            status_code=400,
            detail=f"Invalid {field_name} format, expected YYYY-MM-DD",
        )
    return value


@router.get("/mr-list")
def get_mr_list_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """
    Jute MR List report (header-only).

    Lists Jute MR headers within a date range, each with its status, placed by
    its MR date or, when it has none yet (Open, Pending on a transfer chain),
    its gate-entry date. status_ids narrows to the ticked statuses.
    The Tally sheet (GET /tally-download) stays approved-only whatever is
    shown here.

    Query params:
    - co_id: Company ID (required, int)
    - branch_id: Branch ID (optional, int)
    - date_from: Start date YYYY-MM-DD (required)
    - date_to: End date YYYY-MM-DD (required)
    - status_ids: Comma-separated status ids, e.g. "1,3,13" (optional; empty = all)
    - search: Optional search term (party name, vehicle_no, challan_no,
      formatted MR number)
    - page: Page number (default 1)
    - limit: Page size (default 50, max 10000)

    Returns:
        {"data": [row, ...], "total": <int>}
    """
    try:
        q_co_id = request.query_params.get("co_id")
        q_branch_id = request.query_params.get("branch_id")
        q_date_from = request.query_params.get("date_from")
        q_date_to = request.query_params.get("date_to")
        q_search = request.query_params.get("search")
        q_page = request.query_params.get("page")
        q_limit = request.query_params.get("limit")
        q_status_ids = request.query_params.get("status_ids") or ""

        if not q_co_id:
            raise HTTPException(status_code=400, detail="co_id is required")
        try:
            status_ids = tuple(int(x) for x in q_status_ids.split(",") if x.strip())
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid status_ids")
        if not q_date_from:
            raise HTTPException(status_code=400, detail="date_from is required")
        if not q_date_to:
            raise HTTPException(status_code=400, detail="date_to is required")

        try:
            co_id = int(q_co_id)
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Invalid co_id")

        branch_id = None
        if q_branch_id not in (None, ""):
            try:
                branch_id = int(q_branch_id)
            except (TypeError, ValueError):
                raise HTTPException(status_code=400, detail="Invalid branch_id")

        date_from = _parse_iso_date(q_date_from, "date_from")
        date_to = _parse_iso_date(q_date_to, "date_to")

        try:
            page = int(q_page) if q_page not in (None, "") else 1
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Invalid page")
        try:
            limit = int(q_limit) if q_limit not in (None, "") else 50
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Invalid limit")

        page = max(page, 1)
        limit = max(min(limit, 10000), 1)
        offset = (page - 1) * limit

        search_like = None
        if q_search and q_search.strip():
            search_like = f"%{q_search.strip()}%"

        base_params = {
            "co_id": co_id,
            "branch_id": branch_id,
            "date_from": date_from,
            "date_to": date_to,
            "search_like": search_like,
        }

        list_params = {
            **base_params,
            "limit": limit,
            "offset": offset,
        }

        list_query = get_mr_list_query(status_ids)
        rows = db.execute(list_query, list_params).fetchall()

        data = []
        for row in rows:
            mapped = dict(row._mapping)
            # Convert date-like fields to ISO string for JSON serialization
            for date_field in ("jute_mr_date", "jute_gate_entry_date", "challan_date"):
                value = mapped.get(date_field)
                if value is not None and hasattr(value, "isoformat"):
                    mapped[date_field] = value.isoformat()
            data.append(mapped)

        count_query = get_mr_list_count_query(status_ids)
        count_result = db.execute(count_query, base_params).scalar()
        total = int(count_result) if count_result is not None else 0

        return {"data": data, "total": total}

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute MR list report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute MR list report",
        )


@router.get("/mr-list-statuses")
def get_mr_list_statuses(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """Status options for the MR list report: the statuses the company's MRs carry.

    Query params: co_id (required, int). Returns {"data": [{status_id, status_name}]}.
    """
    try:
        try:
            co_id = int(request.query_params.get("co_id"))
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="co_id is required")
        rows = db.execute(get_mr_list_statuses_query(), {"co_id": co_id}).fetchall()
        return {"data": [dict(r._mapping) for r in rows]}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute MR list statuses: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail="Error fetching jute MR list statuses")


@router.get("/qty-wise")
async def get_jute_qty_wise_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """
    Jute Quantity Wise report (#2) — opening/receipt/issue/closing per jute
    quality over a date range, in weight (kg) and quantity. Also backs the
    Inventory Snapshot (#14): the page passes date_from = FY start, date_to = as-of.

    Query params:
    - branch_id: Branch ID (required)
    - date_from: Start date YYYY-MM-DD (required)
    - date_to: End date YYYY-MM-DD (required)
    """
    try:
        q_branch_id = request.query_params.get("branch_id")
        q_date_from = request.query_params.get("date_from")
        q_date_to = request.query_params.get("date_to")

        if not q_branch_id:
            raise HTTPException(status_code=400, detail="branch_id is required")
        if not q_date_from:
            raise HTTPException(status_code=400, detail="date_from is required")
        if not q_date_to:
            raise HTTPException(status_code=400, detail="date_to is required")

        try:
            branch_id = int(q_branch_id)
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid branch_id")

        date_from = _parse_iso_date(q_date_from, "date_from")
        date_to = _parse_iso_date(q_date_to, "date_to")

        query = get_jute_qty_wise_report_query()
        rows = db.execute(query, {
            "branch_id": branch_id,
            "date_from": date_from,
            "date_to": date_to,
        }).fetchall()

        data = [dict(r._mapping) for r in rows]
        return {"data": data}

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute quantity wise report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute quantity wise report",
        )


@router.get("/txn-summary")
async def get_jute_txn_summary_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """
    Jute Issue & Receipt Summary (#3) — transaction-level rows over a date range.

    Query params:
    - branch_id (required, int)
    - date_from, date_to (required, YYYY-MM-DD)
    - type: 'R' | 'I' | 'ALL' (default 'ALL')
    - search: optional MR-number filter
    """
    try:
        q_branch_id = request.query_params.get("branch_id")
        q_date_from = request.query_params.get("date_from")
        q_date_to = request.query_params.get("date_to")
        q_type = (request.query_params.get("type") or "ALL").upper()
        q_search = request.query_params.get("search")

        if not q_branch_id:
            raise HTTPException(status_code=400, detail="branch_id is required")
        if not q_date_from:
            raise HTTPException(status_code=400, detail="date_from is required")
        if not q_date_to:
            raise HTTPException(status_code=400, detail="date_to is required")
        if q_type not in ("R", "I", "ALL"):
            raise HTTPException(status_code=400, detail="type must be R, I or ALL")

        try:
            branch_id = int(q_branch_id)
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid branch_id")

        date_from = _parse_iso_date(q_date_from, "date_from")
        date_to = _parse_iso_date(q_date_to, "date_to")
        search_like = f"%{q_search.strip()}%" if q_search and q_search.strip() else None

        query = get_jute_txn_summary_query(q_type)
        rows = db.execute(query, {
            "branch_id": branch_id,
            "date_from": date_from,
            "date_to": date_to,
            "search_like": search_like,
        }).fetchall()

        data = []
        for row in rows:
            mapped = dict(row._mapping)
            value = mapped.get("txn_date")
            if value is not None and hasattr(value, "isoformat"):
                mapped["txn_date"] = value.isoformat()
            data.append(mapped)
        return {"data": data}

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute txn summary report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute txn summary report",
        )


@router.get("/period-wise")
async def get_jute_period_wise_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """
    Jute Month Wise (#10) / Day Wise (#11) — receipt vs issue weight & qty
    aggregated per period.

    Query params:
    - branch_id (required, int)
    - date_from, date_to (required, YYYY-MM-DD)
    - period: 'month' (default) | 'day'
    """
    try:
        q_branch_id = request.query_params.get("branch_id")
        q_date_from = request.query_params.get("date_from")
        q_date_to = request.query_params.get("date_to")
        q_period = (request.query_params.get("period") or "month").lower()

        if not q_branch_id:
            raise HTTPException(status_code=400, detail="branch_id is required")
        if not q_date_from:
            raise HTTPException(status_code=400, detail="date_from is required")
        if not q_date_to:
            raise HTTPException(status_code=400, detail="date_to is required")
        if q_period not in ("month", "day"):
            raise HTTPException(status_code=400, detail="period must be month or day")

        try:
            branch_id = int(q_branch_id)
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid branch_id")

        date_from = _parse_iso_date(q_date_from, "date_from")
        date_to = _parse_iso_date(q_date_to, "date_to")

        query = get_jute_period_wise_query(q_period)
        rows = db.execute(query, {
            "branch_id": branch_id,
            "date_from": date_from,
            "date_to": date_to,
        }).fetchall()

        data = [dict(r._mapping) for r in rows]
        return {"data": data}

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute period wise report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute period wise report",
        )


def _require_branch_id(request: Request) -> int:
    """Validate and return the required branch_id query param."""
    q = request.query_params.get("branch_id")
    if not q:
        raise HTTPException(status_code=400, detail="branch_id is required")
    try:
        return int(q)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="Invalid branch_id")


def _require_co_id(request: Request) -> int:
    """Validate and return the required co_id query param."""
    q = request.query_params.get("co_id")
    if not q:
        raise HTTPException(status_code=400, detail="co_id is required")
    try:
        return int(q)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="Invalid co_id")


def _serialize_dates(rows):
    """Map SQLAlchemy rows to dicts, converting any date-like values to ISO."""
    out = []
    for row in rows:
        m = dict(row._mapping)
        for k, v in list(m.items()):
            if v is not None and hasattr(v, "isoformat"):
                m[k] = v.isoformat()
        out.append(m)
    return out


def _branch_daterange_params(request: Request):
    """Validate branch_id + date_from + date_to and return them as a params dict."""
    branch_id = _require_branch_id(request)
    q_from = request.query_params.get("date_from")
    q_to = request.query_params.get("date_to")
    if not q_from:
        raise HTTPException(status_code=400, detail="date_from is required")
    if not q_to:
        raise HTTPException(status_code=400, detail="date_to is required")
    return {
        "branch_id": branch_id,
        "date_from": _parse_iso_date(q_from, "date_from"),
        "date_to": _parse_iso_date(q_to, "date_to"),
    }


@router.get("/percent-claims")
async def get_jute_percent_claims_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """Jute Percent Claims (#7) — per-supplier pass/claim. Params: branch_id, date_from, date_to."""
    try:
        params = _branch_daterange_params(request)
        rows = db.execute(get_jute_percent_claims_query(), params).fetchall()
        return {"data": [dict(r._mapping) for r in rows]}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute percent claims report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute percent claims report",
        )


@router.get("/mukham-moisture")
async def get_jute_mukham_moisture_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """Mukham Moisture (#9) — supplied vs allowed moisture per supplier/mukam. Params: branch_id, date_from, date_to."""
    try:
        params = _branch_daterange_params(request)
        rows = db.execute(get_jute_mukham_moisture_query(), params).fetchall()
        return {"data": [dict(r._mapping) for r in rows]}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute mukham moisture report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute mukham moisture report",
        )


@router.get("/with-value")
async def get_jute_with_value_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """
    Jute with Value (#1) — opening/receipt/issue/closing per quality with
    receipt & issue value and average issue rate, over a date range.

    Query params: branch_id (req), date_from (req), date_to (req).
    """
    try:
        branch_id = _require_branch_id(request)
        q_date_from = request.query_params.get("date_from")
        q_date_to = request.query_params.get("date_to")
        if not q_date_from:
            raise HTTPException(status_code=400, detail="date_from is required")
        if not q_date_to:
            raise HTTPException(status_code=400, detail="date_to is required")
        date_from = _parse_iso_date(q_date_from, "date_from")
        date_to = _parse_iso_date(q_date_to, "date_to")

        rows = db.execute(get_jute_with_value_query(), {
            "branch_id": branch_id,
            "date_from": date_from,
            "date_to": date_to,
        }).fetchall()
        return {"data": [dict(r._mapping) for r in rows]}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute with value report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute with value report",
        )


@router.get("/mr-in-stock")
async def get_jute_mr_in_stock_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """Jute MR in Stock (#4) — per MR-line outstanding stock. Param: branch_id."""
    try:
        branch_id = _require_branch_id(request)
        rows = db.execute(
            get_jute_mr_in_stock_query(), {"branch_id": branch_id}
        ).fetchall()
        return {"data": _serialize_dates(rows)}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute MR in stock report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute MR in stock report",
        )


@router.get("/mr-wise")
async def get_jute_mr_wise_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """MR Wise (#6) — per MR-line received/issued/balance. Param: branch_id."""
    try:
        branch_id = _require_branch_id(request)
        rows = db.execute(
            get_jute_mr_wise_query(), {"branch_id": branch_id}
        ).fetchall()
        return {"data": _serialize_dates(rows)}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute MR wise report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute MR wise report",
        )


@router.get("/godown-wise")
async def get_jute_godown_wise_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """Godown Wise Stock (#5) — balance grouped by warehouse + quality. Param: branch_id."""
    try:
        branch_id = _require_branch_id(request)
        rows = db.execute(
            get_jute_godown_wise_query(), {"branch_id": branch_id}
        ).fetchall()
        return {"data": [dict(r._mapping) for r in rows]}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute godown wise report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute godown wise report",
        )


# The Tally export fills docs/PurchaseFormat.xlsx (Tally Prime "Accounting
# Voucher" sample supplied by the users): headers, widths and the Read Me sheet
# stay exactly as supplied, only the sample rows are replaced.
_PURCHASE_TALLY_TEMPLATE = Path(__file__).resolve().parents[2] / "docs" / "PurchaseFormat.xlsx"
_PURCHASE_TALLY_SHEET = "Accounting Voucher"
_PURCHASE_TALLY_DATE_COLS = ("A", "Q")
# Bill-wise allocation of the supplier ledger, appended after the template's last
# column (Tally field names from its Read Me sheet). Party row only.
_PURCHASE_TALLY_EXTRA_COLUMNS = {"AI": "Bill Type of Ref", "AJ": "Bill Name"}
_PURCHASE_TALLY_MAX_ROWS = 10_000


def _join_text(*parts) -> str:
    return ", ".join(p.strip() for p in parts if p and p.strip())


def build_purchase_tally_rows(headers, lines, tds_for=None) -> list:
    """Build Tally "Accounting Voucher" rows (dicts keyed by column letter).

    headers: one dict per approved MR (get_jute_purchase_tally_header_query).
    lines:   one dict per MR line (get_jute_purchase_tally_line_query).
    tds_for: (header, gross purchase) -> TDS, for MRs that never stored a
             tds_amount (approved before TDS was saved); None = no TDS.

    Per MR, in the order of docs/PurchaseFormat.xlsx: supplier (Cr net),
    purchase ledger (Dr) sharing its row with the first item, the other items,
    claim (Cr), the MR's 194Q TDS (Cr) and round off -- ledger names are the
    company's (header *_ledger columns). Money is built from the lines the way
    the MR does it (compute_jute_totals), so every voucher's Dr equals its Cr.
    MRs without lines, or with a line lacking accepted_weight (pre FY 26-27
    imports kept in quintals), are skipped.
    """
    lines_by_mr = defaultdict(list)
    for ln in lines:
        lines_by_mr[ln["jute_mr_id"]].append(ln)

    rows = []
    for h in headers:
        mr_lines = lines_by_mr.get(h["jute_mr_id"])
        if not mr_lines or any(ln["accepted_weight"] is None for ln in mr_lines):
            continue
        mr_no = h["mr_no"]
        base = {"A": h["invoice_date"] or h["jute_mr_date"], "B": "Purchase", "C": mr_no,
                "N": "Item Invoice"}
        is_bale = (h["unit_conversion"] or "").upper() == "BALE"

        item_rows, claims, total_kg, purchase, claim = [], [], 0.0, 0.0, 0.0
        for ln in mr_lines:
            kg = float(ln["accepted_weight"])
            if kg <= 0:
                continue
            # Rates are keyed per kg and stored per quintal; Tally gets the kg rate at
            # 2 dp and amount = qty x that rate, so its own qty x rate check agrees.
            rate_kg = float((Decimal(str(ln["rate"] or 0)) / 100).quantize(Decimal("0.01"), ROUND_HALF_UP))
            claim_rate = float(ln["claim_rate"] or 0)    # per quintal
            amount = round(kg * rate_kg, 2)
            total_kg += kg
            purchase += amount
            if claim_rate > 0:
                claim += round(kg * claim_rate / 100, 2)
                claims.append(f"(Q): {ln['item_name']}- {_kg_text(kg)} Kg. X Rs. {claim_rate / 100:.2f}/-")
            bales = int(ln["actual_qty"] or 0)
            item_rows.append({
                **base,
                "I": ln["item_name"],
                "J": kg,
                "K": rate_kg,
                "L": "kg",
                "M": amount,
                "O": h["godown_name"] or None,
                "AE": f"Raw Jute: {bales} Bales" if is_bale and bales else "Raw Jute: Loose",
                "AF": mr_no,
            })
        if not item_rows:
            continue

        purchase, claim = round(purchase, 2), round(claim, 2)
        tds = h["tds_amount"]
        if tds is None:
            tds = tds_for(h, purchase) if tds_for else 0
        tds = round(float(tds or 0), 2)
        roundoff, net = compute_jute_totals(purchase, claim, tds)

        party = h["party_name"]
        narration = f"Being {_kg_text(total_kg)} Kg. Raw Jute Purchase From {party}"
        if h["invoice_no"]:
            narration += f" Against Inv: {h['invoice_no']}"
            if h["invoice_date"]:
                narration += f",Dt: {_dmy(h['invoice_date'])}"
        if h["challan_no"]:
            narration += f",CH: {h['challan_no']}"
            if h["challan_date"]:
                narration += f",Dt: {_dmy(h['challan_date'])}"
        net_text = f"{net:.2f}".removesuffix(".00")
        narration += f",MR.NO.{mr_no},Amt Rs. {net_text}/-"

        party_state = (h["party_state"] or "").title() or None
        branch_state = (h["branch_state"] or "").title() or None
        rows.append({
            **base,
            "D": _join_text(h["party_address"], h["party_address2"]) or None,
            "E": h["party_zip"],
            "F": party,
            "G": abs(net),
            "H": "Cr" if net >= 0 else "Dr",
            "P": h["invoice_no"] or None,
            "Q": h["invoice_date"],
            "R": party,
            "S": "India",
            "T": party_state,
            "U": h["party_zip"],
            "V": "Regular" if (h["party_gstin"] or "").strip() else None,
            "W": branch_state,
            "X": h["consignee_name"],
            "Y": h["consignee_name"],
            "Z": "Primary",
            "AA": _join_text(h["branch_address1"], h["branch_address2"]) or None,
            "AB": branch_state,
            "AC": "India",
            "AD": h["branch_zipcode"] or None,
            "AH": narration,
            "AI": "New Ref",
            "AJ": mr_no,
        })
        rows.append({**item_rows[0], "F": h["purchase_ledger"], "G": purchase, "H": "Dr"})
        rows.extend(item_rows[1:])
        if claim > 0:
            rows.append({**base, "F": h["claim_ledger"], "G": claim, "H": "Cr",
                         "AG": " + ".join(claims)})
        if tds:
            rows.append({**base, "F": h["tds_ledger"], "G": abs(tds), "H": "Cr" if tds > 0 else "Dr"})
        if roundoff:
            rows.append({**base, "F": h["round_off_ledger"], "G": abs(roundoff),
                         "H": "Dr" if roundoff > 0 else "Cr"})
    return rows


def _build_purchase_tally_xlsx(rows) -> BytesIO:
    """The user-supplied template with its sample rows replaced by `rows`."""
    wb = load_workbook(_PURCHASE_TALLY_TEMPLATE)
    ws = wb[_PURCHASE_TALLY_SHEET]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for r in [r for r in ws.row_dimensions if r > 1]:
        del ws.row_dimensions[r]                     # delete_rows leaves sample heights behind
    for col, header in _PURCHASE_TALLY_EXTRA_COLUMNS.items():
        ws[f"{col}1"] = header
        ws[f"{col}1"]._style = copy(ws["AH1"]._style)
    columns = [c.column_letter for c in ws[1] if c.value]
    for r, row in enumerate(rows, start=2):
        for col in columns:
            value = row.get(col)
            if value is not None and value != "":
                ws[f"{col}{r}"] = value
        for col in _PURCHASE_TALLY_DATE_COLS:
            ws[f"{col}{r}"].number_format = "d-mmm-yy"
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


@router.get("/tally-download")
def get_jute_tally_download(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """
    Jute Purchase Tally download: docs/PurchaseFormat.xlsx filled with one
    balanced voucher per approved MR (see build_purchase_tally_rows).

    Query params:
      - co_id (required, int): Company ID
      - branch_id (optional, int): Branch filter
      - date_from (required, YYYY-MM-DD): Inclusive start of MR date range
      - date_to (required, YYYY-MM-DD): Inclusive end of MR date range

    Returns the xlsx as an attachment; 400 on bad params or more than
    _PURCHASE_TALLY_MAX_ROWS voucher rows.
    """
    try:
        q_co_id = request.query_params.get("co_id")
        q_branch_id = request.query_params.get("branch_id")
        q_date_from = request.query_params.get("date_from")
        q_date_to = request.query_params.get("date_to")

        if not q_co_id:
            raise HTTPException(status_code=400, detail="co_id is required")
        if not q_date_from:
            raise HTTPException(status_code=400, detail="date_from is required")
        if not q_date_to:
            raise HTTPException(status_code=400, detail="date_to is required")

        try:
            co_id = int(q_co_id)
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Invalid co_id")

        branch_id = None
        if q_branch_id not in (None, ""):
            try:
                branch_id = int(q_branch_id)
            except (TypeError, ValueError):
                raise HTTPException(status_code=400, detail="Invalid branch_id")

        date_from = _parse_iso_date(q_date_from, "date_from")
        date_to = _parse_iso_date(q_date_to, "date_to")

        params = {
            "co_id": co_id,
            "branch_id": branch_id,
            "date_from": date_from,
            "date_to": date_to,
        }
        headers = [dict(r._mapping) for r in
                   db.execute(get_jute_purchase_tally_header_query(), params).fetchall()]
        lines = [dict(r._mapping) for r in
                 db.execute(get_jute_purchase_tally_line_query(), params).fetchall()]

        def tds_for(h, gross):
            # ponytail: one cumulative query per MR without a stored TDS (pre-Sep 2026
            # approvals); fine for a month, slow for a whole FY on a far DB.
            return derive_mr_tds(db, party_id=h["party_id"], mr_date=h["jute_mr_date"],
                                 mr_id=h["jute_mr_id"], branch_mr_no=h["branch_mr_no"],
                                 total_amount=gross)["tds_amount"]

        rows = build_purchase_tally_rows(headers, lines, tds_for)

        if len(rows) > _PURCHASE_TALLY_MAX_ROWS:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Date range too wide — max {_PURCHASE_TALLY_MAX_ROWS} rows, "
                    f"got {len(rows)}"
                ),
            )

        filename = f"jute_purchase_tally_{date_from}_{date_to}.xlsx"
        return StreamingResponse(
            _build_purchase_tally_xlsx(rows),
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
            },
        )

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute tally download: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute tally download",
        )


def _optional_daterange(request: Request):
    """
    Read optional date_from/date_to. Both or neither — a half-open range is
    ambiguous, so supplying only one is a 400.

    Returns (date_from, date_to) or (None, None).
    """
    q_from = request.query_params.get("date_from")
    q_to = request.query_params.get("date_to")
    if not q_from and not q_to:
        return None, None
    if not q_from:
        raise HTTPException(status_code=400, detail="date_from is required")
    if not q_to:
        raise HTTPException(status_code=400, detail="date_to is required")
    return _parse_iso_date(q_from, "date_from"), _parse_iso_date(q_to, "date_to")


def _mr_scoped_params(request: Request):
    """
    Scope shared by every MR-keyed detail report: one company, one branch, and
    then EITHER one picked MR or a receipt-date range.

    A picked MR wins outright — no date filter is built at all, because the user
    asked about that MR and its issues and sales run past any range they typed.

    Returns (params, has_mr_filter, has_date_filter).
    """
    params = {
        "co_id": _require_co_id(request),
        "branch_id": _require_branch_id(request),
    }

    q_mr_id = request.query_params.get("mr_id")
    mr_id = None
    if q_mr_id not in (None, ""):
        try:
            mr_id = int(q_mr_id)
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Invalid mr_id")

    if mr_id is not None:
        params["mr_id"] = mr_id
        return params, True, False

    date_from, date_to = _optional_daterange(request)
    if not date_from:
        raise HTTPException(
            status_code=400,
            detail="date_from is required when mr_id is not supplied",
        )
    params["date_from"] = date_from
    params["date_to"] = date_to
    return params, False, True


@router.get("/mr-options")
def get_jute_mr_options(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """
    MR picker options, shared by the MR-wise Issue Detail and Movement
    Detail reports.

    Query params: co_id (req), branch_id (req), date_from/date_to (optional,
    both or neither) filtered on COALESCE(out_date, jute_mr_date, jute_gate_entry_date).

    Returns {"data": [{jute_mr_id, mr_number, mr_label, mr_date, jute_mr_date,
    out_date, jute_gate_entry_no, party_name}, ...]}, newest first.
    """
    try:
        params = {
            "co_id": _require_co_id(request),
            "branch_id": _require_branch_id(request),
        }
        date_from, date_to = _optional_daterange(request)
        if date_from:
            params["date_from"] = date_from
            params["date_to"] = date_to

        rows = db.execute(
            get_jute_mr_options_query(has_date_filter=date_from is not None),
            params,
        ).fetchall()
        return {"data": _serialize_dates(rows)}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute MR options: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute MR options",
        )


@router.get("/mr-issue-detail")
def get_jute_mr_issue_detail_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """
    MR-wise Issue Detail — every jute issue made against each MR line, flat.
    MR lines with no issues still appear with NULL issue columns.

    Query params:
    - co_id: Company ID (required, int)
    - branch_id: Branch ID (required, int)
    - mr_id: Jute MR ID (optional, int). When given, the DATE RANGE IS IGNORED
      — the user picked one MR and wants all of its issues whatever the period.
    - date_from / date_to: YYYY-MM-DD, required ONLY when mr_id is absent.
      Filtered on COALESCE(out_date, jute_mr_date, jute_gate_entry_date).

    Returns {"data": [...]}.
    """
    try:
        params, has_mr_filter, has_date_filter = _mr_scoped_params(request)
        rows = db.execute(
            get_jute_mr_issue_detail_query(
                has_mr_filter=has_mr_filter,
                has_date_filter=has_date_filter,
            ),
            params,
        ).fetchall()
        return {"data": _serialize_dates(rows)}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute MR issue detail report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute MR issue detail report",
        )


@router.get("/mr-movement-detail")
def get_jute_mr_movement_detail_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """
    MR-wise Movement Detail — the Issue Detail report plus raw-jute sales.

    One row per movement: `movement_type` is 'Issue', 'Sale' or
    'Sale (unattributed)'. An MR line with neither still appears once with a
    NULL movement_type. Rows are never multiplied, so issued_weight and
    sold_weight are both summable.

    Sales carry no line-level MR link, only `sales_invoice_jute.mr_id` at header
    level, so an invoice line is filed under an MR line only when that id
    resolves, the MR sits in the same branch, and the MR line holds the same
    item. Sales failing the guard come back as 'Sale (unattributed)' with the MR
    columns blank and `claimed_mr_id` / `claimed_mr_no` showing what the invoice
    asserted — nothing is silently dropped.

    Query params: identical to /mr-issue-detail.
    - co_id: Company ID (required, int)
    - branch_id: Branch ID (required, int)
    - mr_id: Jute MR ID (optional, int). When given, the DATE RANGE IS IGNORED
      and unattributed sales are suppressed.
    - date_from / date_to: YYYY-MM-DD, required ONLY when mr_id is absent.
      MRs are filtered on receipt date; unattributed sales on invoice_date.

    Returns {"data": [...]}.
    """
    try:
        params, has_mr_filter, has_date_filter = _mr_scoped_params(request)
        rows = db.execute(
            get_jute_mr_movement_detail_query(
                has_mr_filter=has_mr_filter,
                has_date_filter=has_date_filter,
            ),
            params,
        ).fetchall()
        return {"data": _serialize_dates(rows)}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(
            f"Error fetching jute MR movement detail report: {e}", exc_info=True
        )
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute MR movement detail report",
        )


def _po_scoped_params(request: Request):
    """
    Scope for the PO-keyed receipt report: one company, one branch, then
    EITHER one picked PO or a po_date range — the same shape as
    _mr_scoped_params, keyed on po_id.

    Returns (params, has_po_filter, has_date_filter).
    """
    params = {
        "co_id": _require_co_id(request),
        "branch_id": _require_branch_id(request),
    }

    q_po_id = request.query_params.get("po_id")
    po_id = None
    if q_po_id not in (None, ""):
        try:
            po_id = int(q_po_id)
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Invalid po_id")

    if po_id is not None:
        params["po_id"] = po_id
        return params, True, False

    date_from, date_to = _optional_daterange(request)
    if not date_from:
        raise HTTPException(
            status_code=400,
            detail="date_from is required when po_id is not supplied",
        )
    params["date_from"] = date_from
    params["date_to"] = date_to
    return params, False, True


@router.get("/po-options")
def get_jute_po_options(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """
    PO picker options for the PO-wise MR Receipt Detail report.

    Query params: co_id (req), branch_id (req), date_from/date_to (optional,
    both or neither) filtered on po_date.

    Returns {"data": [{jute_po_id, po_number, po_label, po_date, party_name,
    po_weight, po_status}, ...]}, newest first.
    """
    try:
        params = {
            "co_id": _require_co_id(request),
            "branch_id": _require_branch_id(request),
        }
        date_from, date_to = _optional_daterange(request)
        if date_from:
            params["date_from"] = date_from
            params["date_to"] = date_to

        rows = db.execute(
            get_jute_po_options_query(has_date_filter=date_from is not None),
            params,
        ).fetchall()
        return {"data": _serialize_dates(rows)}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute PO options: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute PO options",
        )


@router.get("/po-mr-receipt-detail")
def get_jute_po_mr_receipt_detail_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """
    PO-wise MR Receipt Detail — every MR line received against each PO item,
    flat, for ordered-vs-received comparison (TKT-000115). PO items with no
    receipts still appear with NULL MR columns; MR lines whose item matches
    no PO item come back as row_group=1 with the ordered columns NULL.

    Query params:
    - co_id: Company ID (required, int)
    - branch_id: Branch ID (required, int)
    - po_id: Jute PO ID (optional, int). When given, the DATE RANGE IS IGNORED
      — the user picked one PO and wants all of its receipts.
    - date_from / date_to: YYYY-MM-DD, required ONLY when po_id is absent.
      Filtered on po_date.

    Returns {"data": [...]}.
    """
    try:
        params, has_po_filter, has_date_filter = _po_scoped_params(request)
        rows = db.execute(
            get_jute_po_mr_receipt_detail_query(
                has_po_filter=has_po_filter,
                has_date_filter=has_date_filter,
            ),
            params,
        ).fetchall()
        return {"data": _serialize_dates(rows)}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(
            f"Error fetching jute PO MR receipt detail report: {e}", exc_info=True
        )
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute PO MR receipt detail report",
        )


@router.get("/po-pending")
def get_jute_po_pending_report(
    request: Request,
    db: Session = Depends(get_tenant_db),
    token_data: dict = Depends(get_current_user_with_refresh),
):
    """
    Pending PO report — POs that are still live (not closed / rejected /
    cancelled / draft) with arrival pending against the ordered weight, one
    row per PO item (TKT-000115). Current-position report: no date filter.

    Query params: co_id (required, int), branch_id (required, int).

    Returns {"data": [...]}. The po_* header columns repeat across a PO's
    item rows — only the line columns are summable.
    """
    try:
        params = {
            "co_id": _require_co_id(request),
            "branch_id": _require_branch_id(request),
        }
        rows = db.execute(get_jute_po_pending_query(), params).fetchall()
        return {"data": _serialize_dates(rows)}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching jute pending PO report: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail="Error fetching jute pending PO report",
        )
